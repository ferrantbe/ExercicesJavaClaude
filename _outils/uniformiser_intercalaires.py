#!/usr/bin/env python3
"""Impose le style commun aux titres d'intercalaire d'un deck existant.

    python3 _outils/uniformiser_intercalaires.py deck.pptx [--sortie autre.pptx] [--essai]

Calibri gras, aligne a gauche, 66 pt, un cran plus petit tant que le titre ne
tient pas sur une ligne (regle du 13/10/2026). La regle elle-meme vit dans
_slides/mise_en_forme.py, partagee avec le moteur : un deck regenere et un
deck remanie traite par cet outil ont donc le meme style.

Sert pour les decks deja produits ou remanies a la main ; les decks regeneres
passent par finaliser_deck.py, qui applique la meme regle. Ne touche a rien
d'autre que la zone de titre des intercalaires. Peut etre relance sans risque.
"""
import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_slides"))
from mise_en_forme import uniformiser_titre_intercalaire  # noqa: E402


def titre_intercalaire(slide):
    """Meme reperage que finaliser_deck.py : image pleine page + zone de texte."""
    images = [s for s in slide.shapes if s.shape_type == 13
              and s.width > Inches(12) and s.height > Inches(7)]
    if not images:
        return None
    boites = [s for s in slide.shapes if s.has_text_frame and not s.is_placeholder
              and s.text_frame.text.strip()]
    return min(boites, key=lambda s: (s.top, s.left)) if boites else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--sortie")
    ap.add_argument("--essai", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    for i, slide in enumerate(prs.slides, 1):
        boite = titre_intercalaire(slide)
        if boite is None:
            continue
        avant = {r.font.size.pt if r.font.size else None for p in boite.text_frame.paragraphs for r in p.runs}
        taille = uniformiser_titre_intercalaire(boite, prs.slide_width)
        print(f"  slide {i:2}  {taille:g} pt (avant : {', '.join(str(t) for t in avant)})  {boite.text_frame.text}")
    if not a.essai:
        prs.save(a.sortie or a.deck)


if __name__ == "__main__":
    main()
