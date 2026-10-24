#!/usr/bin/env python3
"""
construire_recueil.py — assemble les quinze fascicules en un seul recueil,
enonces/Recueil_des_exercices.docx (demande du 19/10/2026).

    python3 _outils/construire_recueil.py            (version compacte, retenue)
    python3 _outils/construire_recueil.py --une-page-par-exercice

1. build_recueil.js produit le debut : page de garde, table des matieres,
   presentation commune « Comment travailler ces exercices ».
2. Chaque fascicule (enonces/par_module/XXX_enonces.docx, tel qu'il est commite) y est
   ajoute dans l'ordre des seances, avec docxcompose :
   - sa page de garde est retiree ; son titre devient le titre du module
     (« Module POO — Classes et objets »), en tete de page ;
   - son encadre « Comment travailler » est retire (il est dit une fois au
     debut) ; le tableau de ses exercices reste ;
   - ses titres descendent d'un niveau : exercice en niveau 2, sections en
     niveau 3. La table des matieres reprend les niveaux 1 et 2 : les
     modules, et chaque exercice.
3. Style, typographie et pagination comme pour un fascicule.
4. La table des matieres est remplie : une ligne par module et par
   exercice, cliquable, avec son numero de page lu dans un rendu PDF.
   C'est un vrai champ Word : Word propose de le mettre a jour a
   l'ouverture, et le recalcule avec sa propre mise en page.

Le recueil se reconstruit apres toute modification d'un fascicule.
"""
import copy
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docxcompose.composer import Composer
from lxml import etree

RACINE = Path(__file__).resolve().parent.parent
GEN = RACINE / "_outils/generateurs"
ORDRE = ["POO", "STA", "TMP", "TAB", "ENU", "ABS", "ITF", "PKG", "COL", "EXC",
         "JUN", "GEN", "LAM", "FX", "FIC"]
SORTIE = RACINE / "enonces/Recueil_des_exercices_une_page_par_exercice.docx"
# Variante compacte (essai du 19/10/2026) : les modules commencent toujours
# sur une nouvelle page, mais les exercices d'un module ne sont plus separes
# que par un filet horizontal.
# La version compacte est celle retenue par Frederic (20/10/2026) ; l'autre
# reste disponible avec --une-page-par-exercice.
COMPACT = "--une-page-par-exercice" not in sys.argv
if COMPACT:
    SORTIE = RACINE / "enonces/Recueil_des_exercices.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Marges pour un recueil imprime recto verso et relie a gauche (19/10/2026) :
# marges en miroir, plus larges cote reliure (a gauche sur les pages
# impaires, a droite sur les paires), plus etroites ailleurs. En vingtiemes
# de point : 1 cm = 567.
RELIURE, EXTERIEUR, HAUT, BAS, PIED = 1418, 794, 851, 851, 425   # 2,5 / 1,4 / 1,5 / 1,5 / 0,75 cm
# La Presse imprime en recto seul par defaut (22/10/2026) : la reliure est
# alors toujours a gauche, et les modules commencent sur la page suivante,
# sans page blanche. --recto-verso retrouve les marges en miroir et les
# modules sur une page de droite.
RECTO_VERSO = "--recto-verso" in sys.argv
if not RECTO_VERSO:
    EXTERIEUR = 1134                                               # 2 cm
LARGEUR_TEXTE = 11906 - RELIURE - EXTERIEUR


def texte(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def style_de(p):
    s = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return s.get(qn("w:val")) if s is not None else None


def taille_runs(p):
    return {sz.get(qn("w:val")) for sz in p.iter(qn("w:sz"))}


def filet(p):
    """Remplace le contenu du paragraphe par un filet horizontal."""
    for e in list(p):
        p.remove(e)
    pPr = etree.SubElement(p, qn("w:pPr"))
    bdr = etree.SubElement(pPr, qn("w:pBdr"))
    bas = etree.SubElement(bdr, qn("w:bottom"))
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "A6A6A6")):
        bas.set(qn("w:" + k), v)
    sp = etree.SubElement(pPr, qn("w:spacing"))
    sp.set(qn("w:before"), "240"); sp.set(qn("w:after"), "360")


def preparer(code):
    """Le fascicule, sans sa page de garde ni son encadre, titres decales."""
    d = Document(str(RACINE / f"enonces/par_module/{code}_enonces.docx"))
    corps = d.element.body
    enfants = list(corps)
    # 1. page de garde : jusqu'au paragraphe qui porte le saut de section
    #    (pagination) ou, a defaut, jusqu'au premier saut de page.
    fin = next((i for i, e in enumerate(enfants) if e.tag == qn("w:p") and e.find(".//" + qn("w:sectPr")) is not None), None)
    if fin is None:
        fin = next(i for i, e in enumerate(enfants) if e.tag == qn("w:p") and e.find(".//" + qn("w:br")) is not None
                   and e.find(".//" + qn("w:br")).get(qn("w:type")) == "page")
    garde = enfants[:fin + 1]
    titre = next((texte(p) for p in garde if p.tag == qn("w:p") and "72" in taille_runs(p)), f"Module {code}")
    sous = next((texte(p) for p in garde if p.tag == qn("w:p") and "40" in taille_runs(p)), "")
    for e in garde:
        corps.remove(e)
    # 2. « Comment travailler ces exercices » devient le titre du module ;
    #    l'encadre qui suit est retire.
    enfants = list(corps)
    i = next(i for i, e in enumerate(enfants) if e.tag == qn("w:p") and style_de(e) == "Heading1"
             and "Comment travailler" in texte(e))
    tete = enfants[i]
    runs = tete.findall(qn("w:r"))
    for r in runs[1:]:
        tete.remove(r)
    for t in runs[0].findall(qn("w:t")):
        t.text = ""
    runs[0].find(qn("w:t")).text = f"{titre} — {sous}" if sous else titre
    # le titre du module, plus grand que celui d'un exercice
    for nom in ("w:sz", "w:szCs"):
        e = runs[0].find(qn("w:rPr") + "/" + qn(nom))
        if e is not None:
            e.set(qn("w:val"), "40")
    pPr = tete.find(qn("w:pPr"))
    etree.SubElement(pPr, qn("w:pageBreakBefore"))
    # signet du module, pour la table des matieres
    bm = etree.Element(qn("w:bookmarkStart")); bm.set(qn("w:id"), str(9000 + ORDRE.index(code))); bm.set(qn("w:name"), f"module-{code}")
    be = etree.Element(qn("w:bookmarkEnd")); be.set(qn("w:id"), str(9000 + ORDRE.index(code)))
    pPr.addnext(bm); tete.append(be)
    suivant = enfants[i + 1]
    if suivant.tag == qn("w:tbl"):
        corps.remove(suivant)
    # 3. titres decales d'un niveau (le titre du module reste en niveau 1)
    decalage = {"Heading3": "Heading4", "Heading2": "Heading3", "Heading1": "Heading2"}
    for p in corps.iter(qn("w:p")):
        if p is tete:
            continue
        s = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        if s is not None and s.get(qn("w:val")) in decalage:
            s.set(qn("w:val"), decalage[s.get(qn("w:val"))])
    # 4. variante compacte : chaque saut de page devient un filet
    if COMPACT:
        for p in list(corps.iter(qn("w:p"))):
            brs = p.findall(".//" + qn("w:br"))
            if brs and all(b.get(qn("w:type")) == "page" for b in brs) and not texte(p).strip():
                filet(p)
    # le dernier sectPr (pied de page du fascicule) ne doit pas suivre
    for s in corps.findall(qn("w:sectPr")):
        corps.remove(s)
    return d, (titre, sous)


def marges_reliure(chemin, miroir=None):
    """Marges de reliure dans chaque section ; en miroir dans les reglages
    pour le recto verso seulement."""
    if miroir is None:
        miroir = RECTO_VERSO
    import zipfile, shutil
    with zipfile.ZipFile(chemin) as z:
        parties = {k: z.read(k) for k in z.namelist()}
    doc = parties["word/document.xml"].decode("utf-8")
    doc = re.sub(r"<w:pgMar [^>]*/>",
                 f'<w:pgMar w:top="{HAUT}" w:right="{EXTERIEUR}" w:bottom="{BAS}" w:left="{RELIURE}" '
                 f'w:header="{PIED}" w:footer="{PIED}" w:gutter="0"/>', doc)
    parties["word/document.xml"] = doc.encode("utf-8")
    reg = parties["word/settings.xml"].decode("utf-8")
    if not miroir:
        reg = reg.replace("<w:mirrorMargins/>", "")
    elif "<w:mirrorMargins" not in reg:
        # dans l'ordre du schema : apres zoom et les eventuels reglages d'affichage
        m = re.search(r"<w:zoom[^>]*/>", reg)
        pos = m.end() if m else reg.index(">", reg.index("<w:settings")) + 1
        reg = reg[:pos] + "<w:mirrorMargins/>" + reg[pos:]
    parties["word/settings.xml"] = reg.encode("utf-8")
    with zipfile.ZipFile(str(chemin) + ".tmp", "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in parties.items():
            z.writestr(k, v)
    shutil.move(str(chemin) + ".tmp", chemin)


def chapitres_a_droite(chemin):
    """Chaque module commence sur une page de droite (impaire), comme un
    chapitre de livre (19/10/2026) : un saut de section « page impaire »
    avant chaque module, qui insere une page blanche au besoin.

    Dans Word, le sectPr d'un paragraphe termine une section et dit comment
    ELLE a commence. Le sectPr place avant le module k termine donc la
    section du module k-1, qui commence sur une page impaire ; celui place
    avant le premier module termine l'introduction, qui garde la reprise de
    la numerotation posee par la pagination."""
    import zipfile, shutil
    with zipfile.ZipFile(chemin) as z:
        parties = {k: z.read(k) for k in z.namelist()}
    doc = parties["word/document.xml"].decode("utf-8")
    fin = doc.rindex("<w:sectPr")
    final = doc[fin:doc.index("</w:sectPr>", fin) + len("</w:sectPr>")]
    sans_num = re.sub(r"<w:pgNumType[^>]*/>", "", final)
    def avec_type(sect, typ):
        sect = re.sub(r"<w:type [^>]*/>", "", sect)
        return re.sub(r"(<w:sectPr[^>]*>)", r'\1<w:type w:val="%s"/>' % typ, sect, count=1)
    ancres = [m.start() for m in re.finditer(r'<w:bookmarkStart [^>]*w:name="module-', doc)]
    for k, pos in reversed(list(enumerate(ancres))):
        debut_p = doc.rindex("<w:p>", 0, pos) if doc.rfind("<w:p>", 0, pos) > doc.rfind("<w:p ", 0, pos) else doc.rindex("<w:p ", 0, pos)
        fin_p = doc.index("</w:p>", pos) + len("</w:p>")
        para = doc[debut_p:fin_p].replace("<w:pageBreakBefore/>", "")
        sect = avec_type(final, "nextPage") if k == 0 else avec_type(sans_num, "oddPage")
        doc = doc[:debut_p] + "<w:p><w:pPr>" + sect + "</w:pPr></w:p>" + para + doc[fin_p:]
    fin = doc.rindex("<w:sectPr")
    doc = doc[:fin] + avec_type(sans_num, "oddPage") + doc[doc.index("</w:sectPr>", fin) + len("</w:sectPr>"):]
    parties["word/document.xml"] = doc.encode("utf-8")
    with zipfile.ZipFile(str(chemin) + ".tmp", "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in parties.items():
            z.writestr(k, v)
    shutil.move(str(chemin) + ".tmp", chemin)


def numeroter_comme_imprime(chemin, total):
    """Pour l'impression recto verso, le numero affiche est celui de la page
    physique : la page de garde compte pour 1 (sans numero), la suivante
    porte 2, et une page impaire est toujours une page de droite. Le total
    est celui du document entier (NUMPAGES), page de garde comprise."""
    import zipfile, shutil
    with zipfile.ZipFile(chemin) as z:
        parties = {k: z.read(k) for k in z.namelist()}
    doc = parties["word/document.xml"].decode("utf-8")
    doc = doc.replace('<w:pgNumType w:start="1"/>', '<w:pgNumType w:start="2"/>')
    parties["word/document.xml"] = doc.encode("utf-8")
    for k in parties:
        if re.match(r"word/footer\d+\.xml$", k):
            x = parties[k].decode("utf-8").replace(" SECTIONPAGES ", " NUMPAGES ")
            x = re.sub(r"(NUMPAGES .*?<w:t>)\d+(</w:t>)", lambda m: m.group(1) + str(total) + m.group(2), x, flags=re.S)
            parties[k] = x.encode("utf-8")
    with zipfile.ZipFile(str(chemin) + ".tmp", "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in parties.items():
            z.writestr(k, v)
    shutil.move(str(chemin) + ".tmp", chemin)


def remplir_table(chemin, entrees):
    """Ecrit le resultat du champ TOC : une ligne cliquable par entree."""
    d = Document(str(chemin))
    corps = d.element.body
    instr = next(t for t in corps.iter(qn("w:instrText")) if "TOC" in (t.text or ""))
    p_debut = instr.getparent().getparent()
    # le paragraphe de fin est le suivant qui contient fldChar end
    p = p_debut.getnext()
    while p is not None and not any(f.get(qn("w:fldCharType")) == "end" for f in p.iter(qn("w:fldChar"))):
        nxt = p.getnext(); p.getparent().remove(p); p = nxt
    p_fin = p
    for niveau, ancre, libelle, page in entrees:
        para = etree.SubElement(p_fin.getparent(), qn("w:p"))
        p_fin.addprevious(para)
        pPr = etree.SubElement(para, qn("w:pPr"))
        tabs = etree.SubElement(pPr, qn("w:tabs"))
        tab = etree.SubElement(tabs, qn("w:tab")); tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot"); tab.set(qn("w:pos"), str(LARGEUR_TEXTE))
        sp = etree.SubElement(pPr, qn("w:spacing")); sp.set(qn("w:before"), "120" if niveau == 1 else "0"); sp.set(qn("w:after"), "40")
        ind = etree.SubElement(pPr, qn("w:ind")); ind.set(qn("w:left"), "0" if niveau == 1 else "400")
        lien = etree.SubElement(para, qn("w:hyperlink")); lien.set(qn("w:anchor"), ancre); lien.set(qn("w:history"), "1")
        for morceau in (libelle, "\t" + str(page)):
            r = etree.SubElement(lien, qn("w:r"))
            rPr = etree.SubElement(r, qn("w:rPr"))
            if niveau == 1:
                etree.SubElement(rPr, qn("w:b"))
            if morceau.startswith("\t"):
                etree.SubElement(r, qn("w:tab"))
                t = etree.SubElement(r, qn("w:t")); t.text = morceau[1:]
            else:
                t = etree.SubElement(r, qn("w:t")); t.text = morceau; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    # le debut du champ porte deja « separate » ; le paragraphe de debut garde begin/instr/separate
    d.save(str(chemin))


def pages_des_titres(pdf, titres, repere="Comment travailler ces exercices", etiquette="Recueil"):
    """Numero de page affiche de chaque titre, dans l'ordre du document.

    Un titre est cherche comme une ligne a lui seul (et non au milieu d'une
    phrase qui le cite), et toujours apres le titre precedent : les titres
    se suivent dans l'ordre de la table des matieres."""
    brut = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout.split("\f")
    norm = lambda s: re.sub(r"\s+", " ", s.replace("\u00a0", " ").replace("\u202f", " ")).strip()
    lignes = [[norm(l) for l in p.split("\n") if l.strip()] for p in brut]
    depart = next(i for i, p in enumerate(lignes) if any(repere in l for l in p))
    res, i0 = {}, depart
    for t in titres:
        cle = norm(t)
        trouve = None
        for i in range(i0, len(lignes)):
            if any(l == cle or (len(cle) > 40 and l.startswith(cle[:40])) for l in lignes[i]):
                trouve = i
                break
        if trouve is None:
            continue
        i0 = trouve
        # Le numero est lu dans le pied de page : LibreOffice n'exporte pas
        # les pages blanches inserees avant un module, mais les compte, comme
        # Word, dans la numerotation.
        m = re.search(etiquette + r"\s+(\d+)\s*/", " ".join(lignes[trouve]))
        res[t] = int(m.group(1)) if m else trouve + 1
    return res


def inserer_memento(maitre, composeur):
    """Le memento Exercices, sans sa page de garde, a la place de MEMENTO_ICI :
    c'est la presentation commune du recueil (22/10/2026). Ses titres
    descendent d'un niveau, sous « Comment travailler ces exercices »."""
    memento = Document(str(RACINE / "documents/Memento_Exercices.docx"))
    corps = memento.element.body
    enfants = list(corps)
    fin = next(i for i, e in enumerate(enfants)
               if e.tag == qn("w:p") and e.find(".//" + qn("w:sectPr")) is not None)
    for e in enfants[:fin + 1]:
        corps.remove(e)
    decalage = {"Heading2": "Heading3", "Heading1": "Heading2"}
    for p in corps.iter(qn("w:p")):
        s = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        if s is not None and s.get(qn("w:val")) in decalage:
            s.set(qn("w:val"), decalage[s.get(qn("w:val"))])
    cible = next(e for e in maitre.element.body
                 if e.tag == qn("w:p") and "MEMENTO_ICI" in "".join(x.text or "" for x in e.iter(qn("w:t"))))
    position = list(maitre.element.body).index(cible)
    composeur.insert(position, memento)
    maitre.element.body.remove(cible)


# Le fascicule du fil rouge GameHub vit dans son propre depot, voisin de
# celui-ci ; il est ajoute a la fin du recueil (24/10/2026). La variable
# d'environnement GAMEHUB peut designer un autre emplacement du depot.
GAMEHUB = Path(os.environ.get("GAMEHUB", RACINE.parent / "GameHubClaude")) / "_documents" / "GameHub_enonces.docx"


def preparer_gamehub():
    """Le fascicule GameHub, sans sa page de garde, sous un titre de niveau 1 ;
    ses parties deviennent des titres de niveau 2, avec un signet chacune pour
    la table des matieres."""
    d = Document(str(GAMEHUB))
    corps = d.element.body
    enfants = list(corps)
    fin = next(i for i, e in enumerate(enfants) if e.tag == qn("w:p") and any(
        b.get(qn("w:type")) == "page" for b in e.findall(".//" + qn("w:br"))))
    garde = enfants[:fin + 1]
    textes = [texte(p) for p in garde if p.tag == qn("w:p") and texte(p).strip()]
    titre = "GameHub — " + (textes[2] if len(textes) > 2 else "le fil rouge")
    for e in garde:
        corps.remove(e)
    premier = next(e for e in corps if e.tag == qn("w:p") and style_de(e) == "Heading1")
    tete = copy.deepcopy(premier)
    for r in tete.findall(qn("w:r"))[1:]:
        tete.remove(r)
    r0 = tete.find(qn("w:r"))
    for t in r0.findall(qn("w:t"))[1:]:
        r0.remove(t)
    r0.find(qn("w:t")).text = titre
    for nom in ("w:sz", "w:szCs"):
        e = r0.find(qn("w:rPr") + "/" + qn(nom))
        if e is not None:
            e.set(qn("w:val"), "40")
    premier.addprevious(tete)
    decalage = {"Heading3": "Heading4", "Heading2": "Heading3", "Heading1": "Heading2"}
    entrees, n = [], 0
    for p in list(corps.iter(qn("w:p"))):
        if p is tete:
            continue
        st = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        if st is not None and st.get(qn("w:val")) in decalage:
            st.set(qn("w:val"), decalage[st.get(qn("w:val"))])
            if st.get(qn("w:val")) == "Heading2":
                n += 1
                nom = f"gh-{n}"
                bm = etree.Element(qn("w:bookmarkStart")); bm.set(qn("w:id"), str(9500 + n)); bm.set(qn("w:name"), nom)
                be = etree.Element(qn("w:bookmarkEnd")); be.set(qn("w:id"), str(9500 + n))
                p.find(qn("w:pPr")).addnext(bm); p.append(be)
                entrees.append((2, nom, texte(p)))
    pPr = tete.find(qn("w:pPr"))
    etree.SubElement(pPr, qn("w:pageBreakBefore"))
    bm = etree.Element(qn("w:bookmarkStart")); bm.set(qn("w:id"), "9499"); bm.set(qn("w:name"), "module-GH")
    be = etree.Element(qn("w:bookmarkEnd")); be.set(qn("w:id"), "9499")
    pPr.addnext(bm); tete.append(be)
    if COMPACT:
        for p in list(corps.iter(qn("w:p"))):
            brs = p.findall(".//" + qn("w:br"))
            if brs and all(b.get(qn("w:type")) == "page" for b in brs) and not texte(p).strip():
                filet(p)
    for sp in corps.findall(qn("w:sectPr")):
        corps.remove(sp)
    return d, [(1, "module-GH", titre)] + entrees


def main():
    tmp = Path(tempfile.mkdtemp())
    debut = tmp / "debut.docx"
    subprocess.run(["node", str(GEN / "build_recueil.js"), str(debut)], check=True, capture_output=True)
    maitre = Document(str(debut))
    composeur = Composer(maitre)
    inserer_memento(maitre, composeur)
    entrees = []
    for code in ORDRE:
        d, (titre, sous) = preparer(code)
        entrees.append((1, f"module-{code}", f"{titre} — {sous}" if sous else titre))
        for p in d.element.body.iter(qn("w:p")):
            if style_de(p) == "Heading2":
                bm = p.find(".//" + qn("w:bookmarkStart"))
                if bm is not None:
                    entrees.append((2, bm.get(qn("w:name")), texte(p)))
        composeur.append(d)
    if GAMEHUB.is_file():
        d, ent = preparer_gamehub()
        entrees += ent
        composeur.append(d)
    else:
        print(f"Fascicule GameHub introuvable ({GAMEHUB}) : le recueil est produit sans lui.")
    composeur.save(str(SORTIE))
    outils = RACINE / "_outils"
    subprocess.run([sys.executable, str(outils / "style_documents.py"), str(SORTIE)], check=True, capture_output=True)
    marges_reliure(SORTIE)
    def rendre():
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(SORTIE)], capture_output=True)
        return tmp / (SORTIE.stem + ".pdf")
    def nb_pages(pdf):
        return int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
    # La table des matieres prend d'abord sa place, avec des numeros
    # provisoires ; puis pagination et typographie, comme un fascicule ;
    # enfin les vrais numeros, lus sur le rendu du document termine.
    remplir_table(SORTIE, [(niv, a, l, 0) for niv, a, l in entrees])
    n = nb_pages(rendre())
    subprocess.run([sys.executable, str(outils / "ajouter_pagination.py"), str(SORTIE), "Recueil", str(n - 1)], check=True, capture_output=True)
    if RECTO_VERSO:
        chapitres_a_droite(SORTIE)
    # les renvois aux exercices (« voir COL.1 ») deviennent des liens vers
    # leur enonce (23/10/2026)
    r = subprocess.run([sys.executable, str(outils / "lier_renvois.py"), str(SORTIE)], check=True, capture_output=True, text=True)
    print(r.stdout.strip())
    subprocess.run([sys.executable, str(outils / "typographie.py"), str(SORTIE)], check=True, capture_output=True)
    # Numerotation physique d'abord (la page 2 porte le numero 2), pour que
    # les numeros lus dans les pieds de page soient les numeros definitifs.
    numeroter_comme_imprime(SORTIE, 0)
    for passe in range(2):
        pdf = rendre()
        pages = pages_des_titres(pdf, [l for _, _, l in entrees])
        remplir_table(SORTIE, [(niv, a, l, pages.get(l, "?")) for niv, a, l in entrees])
    # La pagination retire la page vide qui suivait la page de garde : le
    # total se relit sur le document termine.
    pdf = rendre()
    pieds = [int(x) for x in re.findall(r"Recueil\s+(\d+)\s*/", subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout)]
    n = max(pieds) if pieds else nb_pages(pdf)     # pages blanches comprises
    numeroter_comme_imprime(SORTIE, n)
    for f in RACINE.rglob("*.avant"):
        f.unlink()
    manquants = [l for _, _, l in entrees if pages.get(l) is None]
    # la version noir et blanc, pour l'impression (20/10/2026)
    subprocess.run([sys.executable, str(outils / "version_nb.py"), str(SORTIE)], check=True, capture_output=True)
    print(f"Recueil : {len(ORDRE)} modules, {len(entrees) - len(ORDRE)} exercices, {n} pages"
          + (f" ; pages introuvables : {manquants}" if manquants else ""))


if __name__ == "__main__":
    main()
