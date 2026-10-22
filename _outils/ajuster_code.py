#!/usr/bin/env python3
"""Ajuste la largeur des zones de code trop larges pour leur contenu, et les centre.

    python3 _outils/ajuster_code.py deck.pptx [--sortie autre.pptx] [--essai]

La regle (14/10/2026) : une zone de code qui occupe toute la largeur de la slide
ne doit pas laisser plus de 20 % de vide a droite de sa ligne la plus longue.
Sinon, elle est retrecie a la largeur de son contenu, plus une marge de 10 %,
et centree horizontalement. Sa hauteur et sa position verticale ne changent pas.

Ne touche que les zones dont le nom commence par « Code » et qui font au moins
11 pouces de large : une zone deja retrecie a la main, ou placee a cote du
texte, est laissee telle quelle. Une zone que chevauche une autre forme (fleche,
bulle, annotation posee a la main) n'est pas deplacee, pour ne pas desaligner
l'annotation : elle est signalee, a ajuster a la main. Peut etre relance sans risque.
"""
import argparse

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

VIDE_MAX = 0.20          # au-dela, on retrecit
MARGE = 0.10             # vide garde a droite apres ajustement
LARGEUR_MIN = Inches(4)
PLEINE_LARGEUR = Inches(11)
CHASSE_CONSOLAS = 0.55   # largeur d'un caractere Consolas, en em (1126/2048)


def largeur_contenu(zone):
    """Largeur, en EMU, de la ligne la plus longue, marges internes comprises."""
    tf = zone.text_frame
    plus_longue = 0
    for p in tf.paragraphs:
        pouces = 0.0
        for r in p.runs:
            pt = r.font.size.pt if r.font.size else 18
            pouces += len(r.text.expandtabs(4)) * CHASSE_CONSOLAS * pt / 72
        plus_longue = max(plus_longue, pouces)
    body = tf._txBody.find(qn("a:bodyPr"))
    marges = int(body.get("lIns", 91440)) + int(body.get("rIns", 91440))
    return Inches(plus_longue) + marges


def chevauche(a, b):
    return (a.left < b.left + b.width and b.left < a.left + a.width
            and a.top < b.top + b.height and b.top < a.top + a.height)


def annotee(zone, slide):
    titre = slide.shapes.title
    for sh in slide.shapes:
        if sh is zone or sh.is_placeholder or sh.name.startswith(("Repere", "Code")):
            continue
        if titre is not None and sh.shape_id == titre.shape_id:
            continue
        if chevauche(zone, sh):
            return True
    return False


def ajuster_slide(slide, largeur_diapo, signales=None):
    """Rend la liste des (ancienne, nouvelle) largeurs en pouces des zones ajustees."""
    faits = []
    for zone in slide.shapes:
        if not zone.name.startswith("Code") or not zone.has_text_frame or zone.width < PLEINE_LARGEUR:
            continue
        contenu = largeur_contenu(zone)
        if contenu <= 0 or (zone.width - contenu) / zone.width <= VIDE_MAX:
            continue
        if annotee(zone, slide):
            if signales is not None:
                signales.append(zone)
            continue
        nouvelle = max(int(contenu / (1 - MARGE)), LARGEUR_MIN)
        ancienne = zone.width
        zone.width = Emu(nouvelle)
        zone.left = Emu(int((largeur_diapo - nouvelle) / 2))
        faits.append((ancienne / 914400, nouvelle / 914400))
    return faits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--sortie")
    ap.add_argument("--essai", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        signales = []
        for avant, apres in ajuster_slide(slide, prs.slide_width, signales):
            print(f"  slide {i:2} : zone de code {avant:.1f} -> {apres:.1f} pouces, centree")
            total += 1
        for z in signales:
            print(f"  slide {i:2} : zone de code trop large, mais annotee (fleche, bulle...) : a ajuster a la main")
    print(f"{total} zone(s) de code ajustee(s)" + (" (essai, rien d'ecrit)" if a.essai else ""))
    if total and not a.essai:
        prs.save(a.sortie or a.deck)


if __name__ == "__main__":
    main()
