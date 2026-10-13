#!/usr/bin/env python3
"""Fait apparaitre au clic, en fondu, le code place sous le texte d'une slide.

    python3 _outils/animer_code.py deck.pptx [--sortie autre.pptx] [--essai]

La regle (13/10/2026, reprise du deck ENU remanie) : quand une zone de code
se trouve SOUS le texte de la slide, le texte s'affiche d'abord et le code
vient au clic suivant, en fondu. Un code place a cote du texte, ou au-dessus,
reste affiche d'emblee.

Les zones de code sont les formes dont le nom commence par « Code » ; un schema
dessine (groupe nomme « Diagramme... ») suit la meme regle. Une zone
deja animee n'est pas touchee ; une animation existante est gardee, et le
code s'ajoute apres elle. Peut etre relance sans risque.
"""
import argparse
import re

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
TIMING = ('<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
          'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" '
          'dur="indefinite" nodeType="mainSeq"><p:childTnLst/></p:cTn><p:prevCondLst><p:cond evt="onPrev" '
          'delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond '
          'evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
          '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>') % P
FONDU = ('<p:par xmlns:p="%s"><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
         '</p:stCondLst><p:childTnLst><p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/>'
         '</p:stCondLst><p:childTnLst><p:par><p:cTn id="{c}" presetID="10" presetClass="entr" '
         'presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/>'
         '</p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst>'
         '<p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
         '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to>'
         '<p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"><p:cBhvr>'
         '<p:cTn id="{e}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
         '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>') % P


def textes_principaux(slide):
    titre = slide.shapes.title
    return [sh for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
            and not sh.name.startswith(("Code", "Diagramme", "Repere", "Espace réservé du num", "Slide Number"))
            and (titre is None or sh.shape_id != titre.shape_id)]


def sous_le_texte(code, textes):
    """Le code est sous un texte qui le chevauche horizontalement."""
    for t in textes:
        chevauche = min(code.left + code.width, t.left + t.width) - max(code.left, t.left) > Inches(1)
        if chevauche and code.top >= t.top + Inches(0.3) and code.top >= t.top + t.height - Inches(0.6):
            return True
    return False


def animer_slide(slide):
    """Rend le nombre de zones de code animees sur cette slide."""
    codes = sorted([sh for sh in slide.shapes if sh.name.startswith(("Code", "Diagramme"))], key=lambda s: s.top)
    textes = textes_principaux(slide)
    if not codes or not textes:
        return 0
    sld = slide._element
    timing = sld.find(qn("p:timing"))
    deja = set(re.findall(r'spTgt spid="(\d+)"', etree.tostring(timing).decode())) if timing is not None else set()
    a_animer = [c for c in codes if str(c.shape_id) not in deja and sous_le_texte(c, textes)]
    if not a_animer:
        return 0
    if timing is None:
        timing = etree.fromstring(TIMING)
        ext = sld.find(qn("p:extLst"))
        (ext.addprevious(timing) if ext is not None else sld.append(timing))
    seq = next(c for c in timing.iter(qn("p:cTn")) if c.get("nodeType") == "mainSeq")
    liste = seq.find(qn("p:childTnLst"))
    n = max(int(c.get("id")) for c in timing.iter(qn("p:cTn")))
    bld = timing.find(qn("p:bldLst"))
    if bld is None:
        bld = etree.SubElement(timing, qn("p:bldLst"))
    for code in a_animer:
        liste.append(etree.fromstring(FONDU.format(a=n + 1, b=n + 2, c=n + 3, d=n + 4, e=n + 5, spid=code.shape_id)))
        n += 5
        etree.SubElement(bld, qn("p:bldP"), spid=str(code.shape_id), grpId="0", animBg="1")
    return len(a_animer)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--sortie")
    ap.add_argument("--essai", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        n = animer_slide(slide)
        if n:
            print(f"  slide {i:2} : code en fondu au clic")
            total += n
    print(f"{total} zone(s) de code animee(s)" + (" (essai, rien d'ecrit)" if a.essai else ""))
    if not a.essai and total:
        prs.save(a.sortie or a.deck)


if __name__ == "__main__":
    main()
