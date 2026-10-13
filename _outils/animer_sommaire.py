#!/usr/bin/env python3
"""Anime la slide « Au programme » : une section par clic.

    python3 _outils/animer_sommaire.py deck.pptx [--sortie autre.pptx] [--essai]

La regle (13/10/2026) : sur la slide « Au programme », chaque section forme un
groupe (son fond et son texte), et les sections apparaissent au clic, dans
l'ordre de leur numero, balayage depuis la gauche.

Les blocs sont les formes nommees « Sommaire du module » (le fond) et
« Sommaire du module, texte », que le finaliseur pose. La slide est
reconstruite a chaque passage ; une slide deja conforme n'est pas touchee.
"""
import argparse
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
from organiser_exercices import degrouper, poser_animation, _sans_ids  # noqa: E402

TITRE = "Au programme"


def est_bloc(sh):
    return sh.name.startswith("Sommaire du module")


def contient_bloc(sh):
    if sh.shape_type == 6:
        return any(contient_bloc(s) for s in sh.shapes)
    return est_bloc(sh)


def numero(sh):
    m = re.match(r"\s*(\d+)", sh.text_frame.text) if sh.has_text_frame else None
    return int(m.group(1)) if m else 999


def deja_conforme(slide):
    """Chaque section dans son propre groupe, animee au clic dans l'ordre des numeros."""
    groupes = [sh for sh in slide.shapes if sh.shape_type == 6 and contient_bloc(sh)]
    if any(est_bloc(sh) for sh in slide.shapes) or not groupes:
        return False
    for g in groupes:
        textes = [c for c in g.shapes if est_bloc(c) and c.has_text_frame and c.text_frame.text.strip()]
        if len(textes) != 1:
            return False
    ordre = [str(g.shape_id) for g in sorted(groupes, key=lambda g: numero(
        next(c for c in g.shapes if est_bloc(c) and c.has_text_frame and c.text_frame.text.strip())))]
    tim = slide._element.find(qn("p:timing"))
    if tim is None:
        return False
    xml = etree.tostring(tim).decode()
    clics = [c for c in tim.iter(qn("p:cTn")) if c.get("nodeType") == "clickEffect"]
    cibles = [c.find(".//" + qn("p:spTgt")).get("spid") for c in clics]
    return cibles == ordre and "wipe(left)" in xml


def animer_slide(slide):
    """Rend le nombre de sections animees si la slide a change, 0 sinon."""
    t = slide.shapes.title
    if t is None or t.text_frame.text.strip() != TITRE:
        return 0
    if not any(contient_bloc(sh) for sh in slide.shapes):
        return 0
    if deja_conforme(slide):
        return 0

    def etat():
        tim = slide._element.find(qn("p:timing"))
        return etree.tostring(slide.shapes._spTree) + (etree.tostring(tim) if tim is not None else b"")
    avant = etat()
    refaire = True
    while refaire:
        refaire = False
        for sh in list(slide.shapes):
            if sh.shape_type == 6 and contient_bloc(sh):
                degrouper(sh)
                refaire = True
                break
    textes = [sh for sh in slide.shapes if est_bloc(sh) and sh.has_text_frame and sh.text_frame.text.strip()]
    fonds = [sh for sh in slide.shapes if est_bloc(sh) and sh not in textes]
    arbre = slide.shapes._spTree
    sections = []
    for tx in sorted(textes, key=numero):
        fond = min((f for f in fonds if abs(f.top - tx.top) < Inches(0.1) and abs(f.left - tx.left) < Inches(0.6)),
                   key=lambda f: abs(f.left - tx.left), default=None)
        membres = [m for m in (fond, tx) if m is not None]
        membres.sort(key=lambda m: arbre.index(m._element))
        g = slide.shapes.add_group_shape(membres)
        g.name = f"Section {numero(tx)}"
        sections.append(g.shape_id)
    poser_animation(slide, sections)
    return len(sections) if _sans_ids(avant) != _sans_ids(etat()) else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--sortie")
    ap.add_argument("--essai", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        n = animer_slide(slide)
        if n:
            print(f"  slide {i:2} : {n} sections, un clic par section")
            total += n
    if not total:
        print("  rien a faire")
    if not a.essai and total:
        prs.save(a.sortie or a.deck)


if __name__ == "__main__":
    main()
