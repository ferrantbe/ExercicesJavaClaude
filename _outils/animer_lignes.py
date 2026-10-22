#!/usr/bin/env python3
"""Anime les slides faites de lignes sur fond sombre (tableaux) : une ligne par clic.

    python3 _outils/animer_lignes.py deck.pptx [--sortie autre.pptx] [--essai]

La regle (15/10/2026) : quand une slide presente plusieurs lignes sur fond sombre
(un tableau dessine en bandeaux, comme « Qui ecrit quoi » ou « Laquelle
choisir ? »), chaque ligne forme un groupe (son bandeau et ses textes) et
apparait au clic, balayage depuis la gauche, de haut en bas. La ligne d'en-tete
et une eventuelle phrase d'introduction restent affichees d'emblee.

Une ligne = un bandeau (rectangle arrondi sans texte, fond 1D1F22, au moins
6 pouces de large) et les formes posees dessus. Ne touche pas les slides
d'exercices, ni « Au programme » (elles ont leurs propres regles), ni une slide
qui porte deja d'autres animations. Une slide deja animee ligne par ligne,
meme autrement (balayage depuis le bas, par exemple), est laissee telle quelle.
"""
import argparse
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
from organiser_exercices import degrouper, poser_animation  # noqa: E402

FOND = "1D1F22"


def est_bandeau(sh):
    if sh.shape_type != 1 or sh.width < Inches(6):
        return False
    if sh.has_text_frame and sh.text_frame.text.strip():
        return False
    try:
        return sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == FOND
    except Exception:
        return False


def contient_bandeau(sh):
    if sh.shape_type == 6:
        return any(contient_bandeau(c) for c in sh.shapes)
    return est_bandeau(sh)


def bandeaux(sh):
    if sh.shape_type == 6:
        return [b for c in sh.shapes for b in bandeaux(c)]
    return [sh] if est_bandeau(sh) else []


def cibles_animees(slide):
    tim = slide._element.find(qn("p:timing"))
    if tim is None:
        return set()
    return {e.get("spid") for e in tim.iter(qn("p:spTgt"))}


def animer_slide(slide):
    """Rend le nombre de lignes animees si la slide a change, 0 sinon (ou un message)."""
    titre = slide.shapes.title
    t = titre.text_frame.text.strip() if titre is not None else ""
    if t.startswith(("Exercice", "Le fil rouge", "Au programme")):
        return 0
    tous = [b for sh in slide.shapes for b in bandeaux(sh)]
    if len(tous) < 2:
        return 0
    cibles = cibles_animees(slide)
    groupes = [sh for sh in slide.shapes if sh.shape_type == 6 and contient_bandeau(sh)]
    # deja conforme : une ligne par groupe, et chaque groupe anime
    if (groupes and all(len(bandeaux(g)) == 1 for g in groupes)
            and len(groupes) == len(tous) and {str(g.shape_id) for g in groupes} <= cibles):
        return 0
    if cibles:
        return "animations existantes : laissee telle quelle"
    # tout a plat, puis un groupe par ligne
    refaire = True
    while refaire:
        refaire = False
        for sh in list(slide.shapes):
            if sh.shape_type == 6 and contient_bandeau(sh):
                degrouper(sh)
                refaire = True
                break
    lignes = sorted([sh for sh in slide.shapes if est_bandeau(sh)], key=lambda b: b.top)
    arbre = slide.shapes._spTree
    ids = []
    for k, b in enumerate(lignes, 1):
        membres = [b] + [sh for sh in slide.shapes
                         if sh is not b and not sh.is_placeholder and sh.shape_type != 6
                         and not sh.name.startswith("Repere")
                         and sh.top >= b.top - Inches(0.05)
                         and sh.top + sh.height <= b.top + b.height + Inches(0.1)
                         and sh.left >= b.left - Inches(0.05)
                         and sh.left + sh.width <= b.left + b.width + Inches(0.1)]
        membres.sort(key=lambda m: arbre.index(m._element))
        g = slide.shapes.add_group_shape(membres)
        g.name = f"Ligne {k}"
        ids.append(g.shape_id)
    poser_animation(slide, ids)
    return len(ids)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--sortie")
    ap.add_argument("--essai", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        n = animer_slide(slide)
        if isinstance(n, str):
            print(f"  slide {i:2} : {n}")
        elif n:
            print(f"  slide {i:2} : {n} lignes, une par clic")
            total += 1
    if not total:
        print("  rien a faire")
    if total and not a.essai:
        prs.save(a.sortie or a.deck)


if __name__ == "__main__":
    main()
