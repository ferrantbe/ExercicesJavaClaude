#!/usr/bin/env python3
"""Configuration du deck FX3 pour construire_deck.py et finaliser_deck.py.

FX3 est un code de DECK, pas de theme : la calculatrice de FX.9, construite
pas a pas. Il reprend la progression de l'ancien deck 09 (slides 24 a 74),
reecrite pour le projet Maven du cours et alignee sur les noms du corrige
de FX.9 (Controleur, Modele, affichage, chiffrePresse, operateurPresse,
effacer, calculer). Il se joue apres FX2, en accompagnement de FX.9.

Usage :
    python3 construire_deck.py module_fx3 FX3_deck.pptx
    python3 finaliser_deck.py FX3 FX3_deck.pptx FX3_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent

CODE = "FX3"
ORDRE = 14   # rang du deck dans le cours : partage avec FX1, decide
SCENARIO = ICI / "scenario" / "Scenario_FX3.md"
GABARIT = ICI / "sources" / "gabarit_deck.pptx"
SOURCES = {"d09": ICI / "sources" / "d09.pptx"}

# Slides prelevees dans les decks sources : (deck, numero[, nom de forme])
GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d09", 7)        # « JavaFX », meme structure que d10/2
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"FX": "fascicule FX", "G": "fil rouge GameHub"}

# Lu par finaliser_deck.py.
# SECTIONS : le texte des intercalaires du deck, numero retire, mot pour mot
#            (casse mise a part) ; un intercalaire absent de cette liste
#            n'est ni numerote ni repere.
# LIBELLES : la version courte, pour le sommaire et le repere en bas de slide.
SECTIONS = [
    "Ce qu'on va construire",
    "La vue, pas à pas",
    "Le contrôleur",
    "Le modèle, et la fin",
]

LIBELLES = [
    "Le projet",
    "La vue",
    "Le contrôleur",
    "Le modèle",
]

# Titre des intercalaires, sans numero : le finaliseur le pose.
# Les titres sont ceux de SECTIONS, dans le meme ordre : ils en sont tires,
# pour qu'un titre change a un seul endroit.
INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}

# Une seule image pour les quatre sections : celle de l'ancien deck, deja
# dans la slide modele. Frederic les differenciera ensuite dans PowerPoint.
IMAGES_INTERCALAIRE = {}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "new", None), ("1.3", "new", None), ("1.4", "new", None),
    ("1.5", "table", None),

    ("2.1", "inter", None),
    ("2.2", "new", None), ("2.3", "new", None), ("2.4", "new", None),
    ("2.5", "new", None), ("2.6", "new", None), ("2.7", "new", None),
    ("2.8", "new", None), ("2.9", "new", None), ("2.10", "new", None),
    ("2.11", "new", None),

    ("3.1", "inter", None),
    ("3.2", "new", None), ("3.3", "new", None), ("3.4", "new", None),
    ("3.5", "new", None), ("3.6", "new", None),

    ("4.1", "inter", None),
    ("4.2", "new", None), ("4.3", "new", None), ("4.4", "new", None),
    ("4.5", "new", None), ("4.6", "new", None), ("4.7", "new", None),
    ("4.8", "renvoi", None),
]

# L'intercalaire du deck 09 porte son titre en 60 pt, comme celui du deck
# 10 : meme taille reduite que FX1, pour que les deux decks se ressemblent.
TAILLE_INTERCALAIRE = 44


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
