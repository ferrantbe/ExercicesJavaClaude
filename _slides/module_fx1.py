#!/usr/bin/env python3
"""Configuration du deck FX1 pour construire_deck.py et finaliser_deck.py.

FX1 est un code de DECK, pas de theme : le theme FX est decoupe en deux
decks, FX1 (FX.1 a FX.5, les evenements) et FX2 (FX.6 a FX.9, proprietes
et architecture). Les exercices gardent leurs codes FX.n, dans un seul
fascicule.

Arborescence attendue, relative au dossier de ce fichier :
    sources/gabarit_deck.pptx  gabarit commun : theme, dispositions, 4 slides modeles
    sources/d10.pptx           ancien deck 10, qui fournit l'intercalaire
    scenario/Scenario_FX1.md   le scenario

La matiere vient de sources/d10.pptx et du debut de sources/d09.pptx,
reecrite dans le scenario : aucune slide n'en est importee telle quelle.

Usage :
    python3 construire_deck.py module_fx1 FX1_deck.pptx
    python3 finaliser_deck.py FX1 FX1_deck.pptx FX1_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent

CODE = "FX1"
ORDRE = 14   # rang du deck dans le cours : partage avec FX2, decide
SCENARIO = ICI / "scenario" / "Scenario_FX1.md"
GABARIT = ICI / "sources" / "gabarit_deck.pptx"
SOURCES = {"d10": ICI / "sources" / "d10.pptx"}

# Slides prelevees dans les decks sources : (deck, numero[, nom de forme])
GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d10", 2)        # « Les evenements en JavaFX », meme structure que d06/20
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"FX": "fascicule FX", "G": "fil rouge GameHub"}

# Lu par finaliser_deck.py.
# SECTIONS : le texte des intercalaires du deck, numero retire, mot pour mot
#            (casse mise a part) ; un intercalaire absent de cette liste
#            n'est ni numerote ni repere.
# LIBELLES : la version courte, pour le sommaire et le repere en bas de slide.
SECTIONS = [
    "Lancer une application JavaFX",
    "Une fenêtre et ses composants",
    "Brancher un bouton",
    "Ce que dit un événement",
    "Un événement voyage",
]

LIBELLES = [
    "Le lancement",
    "La fenêtre",
    "Les gestionnaires",
    "Les événements",
    "La propagation",
]

# Titre des intercalaires, sans numero : le finaliseur le pose.
# Les titres sont ceux de SECTIONS, dans le meme ordre : ils en sont tires,
# pour qu'un titre change a un seul endroit.
INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}

# Une seule image pour les cinq sections : celle de l'ancien deck, deja
# dans la slide modele. Frederic les differenciera ensuite dans PowerPoint.
IMAGES_INTERCALAIRE = {}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "new", None), ("1.3", "new", None), ("1.4", "new", None),

    ("2.1", "inter", None),
    ("2.2", "new", None), ("2.3", "new", None), ("2.4", "table", None),
    ("2.5", "new", None),

    ("3.1", "inter", None),
    ("3.2", "new", None), ("3.3", "demo", "Demo_FX"),
    ("3.4", "new", None), ("3.5", "new", None), ("3.6", "renvoi", None),

    ("4.1", "inter", None),
    ("4.2", "demo", "Demo_FX"),
    ("4.3", "new", None), ("4.4", "table", None), ("4.5", "renvoi", None),
    ("4.6", "new", None), ("4.7", "new", None), ("4.8", "new", None),

    ("5.1", "inter", None),
    ("5.2", "demo", "Demo_FX"),
    ("5.3", "new", None), ("5.4", "table", None), ("5.5", "new", None),
    ("5.6", "renvoi", None),

    ("6.1", "renvoi", None),
]

# L'intercalaire du deck 10 porte son titre en 60 pt : le plus long des
# titres (« 1. Lancer une application JavaFX », 32 caracteres) passait a
# la ligne. Reglee sur un premier rendu.
TAILLE_INTERCALAIRE = 44


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
