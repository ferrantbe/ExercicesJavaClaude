#!/usr/bin/env python3
"""Moteur de construction du deck d'un module, commun a tous les modules.

Extrait de construire_deck_poo.py, dont il reprend le lecteur de scenario,
le texte riche, la coloration Java, les zones de code, les slides de renvoi
et de demonstration. Il reprend aussi le moteur de mise en page de
retouches_1609.py : les puces visent 28 pt et le code 18 pt, ce qui etait
auparavant applique apres coup.

Ce qui est propre a un module vit dans un fichier module_xxx.py, passe en
argument. Voir module_tmp.py pour le modele.

Usage :
    python3 construire_deck.py module_tmp [sortie.pptx]

Dependances : python-pptx, lxml
"""
import copy
import importlib
import io
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mise_en_forme import cran_dessous, uniformiser_titre_intercalaire  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_outils"))
from organiser_exercices import LOGO_GAMEHUB, traiter_slide as organiser_exercices  # noqa: E402
from pptx.util import Emu, Inches, Pt

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

LARGEUR, HAUTEUR = 13.333, 7.5
GAUCHE, DROITE = 0.7, 12.3
HAUT, BAS_PUCES, BAS_CODE = 1.5, 6.75, 6.92

FOND_CODE = RGBColor(0x1D, 0x1F, 0x22)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
COULEURS = {
    "texte": RGBColor(0xBC, 0xBE, 0xC4),
    "mot": RGBColor(0xCF, 0x8E, 0x6D),
    "chaine": RGBColor(0x6A, 0xAB, 0x73),
    "commentaire": RGBColor(0x7A, 0x7E, 0x85),
    "nombre": RGBColor(0x2A, 0xAC, 0xB8),
    "annotation": RGBColor(0xB3, 0xAE, 0x60),
    "methode": RGBColor(0x56, 0xA8, 0xF5),
}


# ---------------------------------------------------------------------------
#  Lecture du scenario
# ---------------------------------------------------------------------------
# Rubriques reconnues : **Source**, **Slide**, **Note**, **Reserve**.
# Seules **Slide** et **Note** produisent quelque chose ; les deux autres
# sont des commentaires de fabrication.

RUBRIQUES = {"**Source**": "source", "**Slide**": "slide",
             "**Note**": "note", "**Réserve**": "reserve",
             "**Reserve**": "reserve"}


def lire_scenario(chemin):
    texte = Path(chemin).read_text(encoding="utf-8")
    blocs = {}
    for morceau in re.split(r"^### ", texte, flags=re.M)[1:]:
        morceau = re.split(r"^## ", morceau, flags=re.M)[0]
        entete, _, corps = morceau.partition("\n")
        ident, _, titre = entete.partition(" · ")
        titre = re.sub(r"\s*\*\(.*?\)\*", "", titre).strip()
        bloc = {"titre": titre, "puces": [], "code": [], "note": [], "table": [], "intro": []}
        section, dans_code, tampon, paragraphe = None, False, [], []
        for ligne in corps.splitlines():
            if dans_code:
                if ligne.startswith("```"):
                    bloc["code"].append("\n".join(tampon))
                    tampon, dans_code = [], False
                else:
                    tampon.append(ligne)
                continue
            if ligne.startswith("```"):
                dans_code = True
                continue
            rubrique = next((v for k, v in RUBRIQUES.items() if ligne.startswith(k)), None)
            if rubrique:
                if paragraphe:
                    bloc["note"].append(" ".join(paragraphe))
                    paragraphe = []
                section = rubrique
                continue
            if ligne.strip() == "---":
                section = None
                continue
            if section == "slide":
                if ligne.startswith("- "):
                    bloc["puces"].append(ligne[2:].strip())
                elif ligne.startswith("|") and not re.match(r"^\|[-| :]+\|$", ligne):
                    bloc["table"].append([c.strip() for c in ligne.strip("|").split("|")])
                elif ligne.strip() and not ligne.startswith(("|", "*", "#")) and not bloc["table"]:
                    # une phrase d'introduction, avant un tableau (15/10/2026)
                    bloc["intro"].append(ligne.strip())
            elif section == "note":
                if ligne.strip():
                    paragraphe.append(ligne.strip())
                elif paragraphe:
                    bloc["note"].append(" ".join(paragraphe))
                    paragraphe = []
        if paragraphe:
            bloc["note"].append(" ".join(paragraphe))
        if bloc["note"] and bloc["note"][0].startswith("*(aucune"):
            bloc["note"] = []
        blocs[ident.strip()] = bloc
    return blocs


# ---------------------------------------------------------------------------
#  Texte riche : `code` et **gras**
# ---------------------------------------------------------------------------

JETONS = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*\([^)]*\)\*)")

# Consolas est plus large que Corbel a taille egale : dans du texte courant,
# un mot de code se pose quatre points plus bas que la prose qui l'entoure
# (28 pt de prose -> 24 pt de code). Les deux occupent alors la meme largeur
# par caractere, ce qui laisse les mesures de hauteur valables.


def taille_inline(t):
    """Taille d'un mot de code dans de la prose de taille t : un cran en dessous."""
    return cran_dessous(t)


def est_code(run):
    return (run.font.name or "") == "Consolas"


def segments(texte):
    for part in JETONS.split(texte):
        if not part:
            continue
        if part.startswith("`"):
            yield part[1:-1], "code"
        elif part.startswith("**"):
            yield part[2:-2], "gras"
        elif part.startswith("*("):
            yield part[1:-1], "italique"
        else:
            yield part, None


def nu(texte):
    """Le texte sans son balisage, pour mesurer sa longueur."""
    return re.sub(r"[`*]", "", texte)


def remplir_paragraphe(par, texte, taille=None, couleur_code=True):
    for morceau, style in segments(texte):
        run = par.add_run()
        run.text = morceau
        if taille:
            run.font.size = Pt(taille_inline(taille.pt)) if style == "code" else taille
        if style == "code":
            run.font.name = "Consolas"
            if couleur_code:
                run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
        elif style == "gras":
            run.font.bold = True
        elif style == "italique":
            run.font.italic = True


# ---------------------------------------------------------------------------
#  Coloration du code
# ---------------------------------------------------------------------------

MOTS = set("""abstract boolean break byte case catch char class continue default do
double else enum extends final finally float for if implements import instanceof int
interface long new null package private protected public return short static super
switch this throw throws true false try void while var record""".split())
LEXEME = re.compile(
    r'(?P<commentaire>//.*$)|(?P<chaine>"(?:[^"\\]|\\.)*"|\'(?:[^\'\\]|\\.)\')'
    r'|(?P<annotation>@\w+)|(?P<nombre>\b\d+(?:\.\d+)?\b)'
    r'|(?P<ident>\b[A-Za-z_]\w*\b)(?P<appel>\s*\()?|(?P<autre>.)')


def lexemes(ligne):
    for m in LEXEME.finditer(ligne):
        genre = m.lastgroup
        if m.group("ident"):
            mot = m.group("ident")
            if mot in MOTS:
                yield mot, "mot"
            elif m.group("appel"):
                yield mot, "methode"
            else:
                yield mot, "texte"
            if m.group("appel"):
                yield m.group("appel"), "texte"
        else:
            yield m.group(0), genre if genre in COULEURS else "texte"


def zone_code(slide, code, x, y, largeur, hauteur):
    boite = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(largeur), Inches(hauteur))
    boite.name = "Code"
    boite.fill.solid()
    boite.fill.fore_color.rgb = FOND_CODE
    tf = boite.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    for i, ligne in enumerate(code.split("\n")):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_before = par.space_after = Pt(0)
        par.line_spacing = 1.0
        ligne = ligne.expandtabs(4) or " "
        for morceau, genre in lexemes(ligne):
            run = par.add_run()
            run.text = morceau
            run.font.name = "Consolas"
            run.font.size = Pt(16)
            run.font.color.rgb = COULEURS[genre]
    return boite


# ---------------------------------------------------------------------------
#  Mise en page : puces visees a 28 pt, code vise a 18 pt
#  (repris de retouches_1609.py, ou il vivait comme retouche apres coup)
# ---------------------------------------------------------------------------

def haut_puces(textes, largeur, t):
    car = max(1, int((largeur - 0.45) * 72 / (0.55 * t)))
    lignes = sum(max(1, -(-len(nu(x)) // car)) for x in textes)
    return lignes * t * 1.2 / 72 + len(textes) * 10 / 72 + 0.15


def taille_puces(textes, largeur, hauteur):
    for t in (28, 26, 24, 22, 20):
        if haut_puces(textes, largeur, t) <= hauteur:
            return t
    return 20


def mesure_code(code):
    lignes = [l.expandtabs(4) for l in code.split("\n")]
    return len(lignes), max((len(l) for l in lignes), default=1) or 1


LARGEUR_CAR = 0.62   # largeur d'un caractere Consolas, en em, avec marge


def taille_code(n, long_max, largeur, hauteur, maxi=18.0):
    tw = (largeur - 0.35) * 72 / (LARGEUR_CAR * long_max)
    th = (hauteur - 0.25) * 72 / (1.2 * n)
    return max(10.0, min(maxi, int(min(tw, th) * 2) / 2))


def valeur(c, t):
    """Code vise 18 pt, puces visees 28 pt ; sous 14 / 22, on evite."""
    return c + 0.5 * t - (0 if c >= 14 else 20) - (0 if t >= 22 else 20)


def poser_taille(forme, t, code_a_part=True):
    """Pose t sur la prose ; les mots de code descendent a taille_inline(t).

    code_a_part=False pour une zone entierement en code, ou la taille voulue
    est celle qu'on passe.
    """
    for p in forme.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(taille_inline(t) if code_a_part and est_code(r) else t)


def disposer(corps, code_boite, textes, code):
    """Choisit entre code a droite et code dessous, puis pose tailles et boites."""
    n, lm = mesure_code(code)
    # Les candidats "dessous" passent en premier : a note egale, max retient le
    # premier, et le code pleine largeur vaut mieux qu'une colonne de puces etroite.
    candidats = []
    for t in (28, 26, 24):
        h = haut_puces(textes, 11.9, t)
        y = HAUT + h + 0.15
        c = taille_code(n, lm, 11.9, BAS_PUCES - y)
        candidats.append((valeur(c, t), "dessous", h, t, c))
    for lb in (5.6, 5.2, 4.8):
        t = taille_puces(textes, lb, BAS_PUCES - HAUT)
        x = GAUCHE + lb + 0.25
        c = taille_code(n, lm, DROITE - x, BAS_CODE - HAUT)
        candidats.append((valeur(c, t), "droite", lb, t, c))
    _, mode, dim, t, c = max(candidats, key=lambda k: k[0])
    poser_taille(corps, t)
    poser_taille(code_boite, c, code_a_part=False)
    hauteur_code = n * c * 1.2 / 72 + 0.3
    corps.left, corps.top = Inches(GAUCHE), Inches(HAUT)
    if mode == "droite":
        corps.width, corps.height = Inches(dim), Inches(BAS_PUCES - HAUT)
        x = GAUCHE + dim + 0.25
        largeur = min(DROITE - x, lm * LARGEUR_CAR * c / 72 + 0.45)
        code_boite.left, code_boite.top = Inches(x), Inches(HAUT)
        code_boite.width = Inches(largeur)
        code_boite.height = Inches(min(hauteur_code, BAS_CODE - HAUT))
    else:
        corps.width, corps.height = Inches(11.9), Inches(dim)
        y = HAUT + dim + 0.15
        code_boite.left, code_boite.top = Inches(GAUCHE), Inches(y)
        code_boite.width = Inches(11.9)
        code_boite.height = Inches(min(hauteur_code, BAS_PUCES - y))
    return mode, t, c


# ---------------------------------------------------------------------------
#  Manipulation des slides
# ---------------------------------------------------------------------------

def layout(prs, nom):
    for l in prs.slide_layouts:
        if l.name == nom:
            return l
    raise KeyError(nom)


def vider_placeholders(slide):
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)


def copier_formes(source, cible):
    """Copie toutes les formes d'une slide dans une autre, images comprises."""
    arbre = cible.shapes._spTree
    for el in source.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        nouveau = copy.deepcopy(el)
        for noeud in nouveau.iter():
            for attr in list(noeud.attrib):
                if not attr.startswith("{%s}" % NS_R):
                    continue
                rid = noeud.get(attr)
                rel = source.part.rels.get(rid)
                if rel is None:
                    del noeud.attrib[attr]
                elif rel.is_external:
                    noeud.set(attr, cible.part.relate_to(rel.target_ref, rel.reltype,
                                                         is_external=True))
                elif "image" in rel.reltype:
                    _, nid = cible.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                    noeud.set(attr, nid)
                else:
                    del noeud.attrib[attr]
        arbre.append(nouveau)


def importer_slide(prs, source):
    nouvelle = prs.slides.add_slide(layout(prs, source.slide_layout.name))
    vider_placeholders(nouvelle)
    copier_formes(source, nouvelle)
    return nouvelle


def sans_puce(par):
    """Retire la puce et le retrait du paragraphe : le texte part de la marge."""
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    for enfant in pPr.findall(qn("a:buNone")):
        pPr.remove(enfant)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def ecrire_notes(slide, paragraphes):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    for i, texte in enumerate(paragraphes):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        remplir_paragraphe(par, texte, couleur_code=False)


def ajouter_numero(slide, modele):
    if modele is not None:
        slide.shapes._spTree.append(copy.deepcopy(modele))


TAILLE_TITRE = 40


def titre(slide, texte):
    """Le titre accepte le meme balisage que le corps : `code`, **gras**.

    Un mot de code y est en Consolas jaune, a taille_inline de la taille du
    titre - meme regle que dans le texte courant.
    """
    tf = slide.shapes.title.text_frame
    tf.clear()
    t = TAILLE_TITRE if len(nu(texte)) <= 34 else 34
    remplir_paragraphe(tf.paragraphs[0], texte, Pt(t))


def poser_une_ligne(boite, texte):
    tf = boite.text_frame
    premier = tf.paragraphs[0]
    premier.runs[0].text = texte
    for r in premier.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def remplacer_fond(slide, chemin):
    """Echange l'image pleine page d'un intercalaire contre une autre.

    La nouvelle image reprend la position et la taille de l'ancienne, puis
    passe en arriere-plan pour que le bandeau de titre reste lisible.
    """
    images = [f for f in slide.shapes if f.shape_type == 13]
    if not images:
        return
    fond = max(images, key=lambda f: f.width * f.height)
    boite = (fond.left, fond.top, fond.width, fond.height)
    fond._element.getparent().remove(fond._element)
    pic = slide.shapes.add_picture(str(chemin), *boite)
    arbre = slide.shapes._spTree
    arbre.remove(pic._element)
    arbre.insert(2, pic._element)   # apres nvGrpSpPr et grpSpPr


def textes_intercalaire(slide, titre_section, sous_titre):
    boites = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()
              and not s.is_placeholder]
    boites.sort(key=lambda s: (s.top, s.left))
    poser_une_ligne(boites[0], titre_section)
    # style commun a tous les decks (mise_en_forme.py), quel que soit le deck source
    uniformiser_titre_intercalaire(boites[0], Inches(LARGEUR))
    if len(boites) > 1:
        if sous_titre:
            poser_une_ligne(boites[1], sous_titre)
        else:
            boites[1]._element.getparent().remove(boites[1]._element)


# ---------------------------------------------------------------------------
#  Fabriques de slides
# ---------------------------------------------------------------------------

def puces_visibles(bloc):
    return [p for p in bloc["puces"] if not p.startswith("*(")]


def slide_contenu(prs, bloc, numero):
    slide = prs.slides.add_slide(layout(prs, "Titre et contenu"))
    titre(slide, bloc["titre"])
    corps = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1][0]
    puces = puces_visibles(bloc)
    code = bloc["code"][0] if bloc["code"] else None
    tf = corps.text_frame
    tf.clear()
    for i, texte in enumerate(puces):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(10)
        sans_puce(par)
        remplir_paragraphe(par, texte)
    if code:
        boite = zone_code(slide, code, GAUCHE, HAUT, 6.0, BAS_CODE - HAUT)
        disposer(corps, boite, [nu(p) for p in puces], code)
    else:
        corps.left, corps.top = Inches(GAUCHE), Inches(HAUT)
        corps.width, corps.height = Inches(11.9), Inches(BAS_PUCES - HAUT)
        poser_taille(corps, taille_puces([nu(p) for p in puces], 11.9, BAS_PUCES - HAUT))
    ajouter_numero(slide, numero)
    return slide


def slide_tableau(prs, bloc, numero):
    """Une slide dont le contenu est le tableau markdown du scenario."""
    slide = prs.slides.add_slide(layout(prs, "Titre seul"))
    titre(slide, bloc["titre"])
    lignes = bloc["table"]
    entete, corps = lignes[0], lignes[1:]
    colonnes = len(entete)
    largeurs = []
    for c in range(colonnes):
        long_max = max(len(nu(l[c])) for l in lignes)
        largeurs.append(max(1.6, long_max * 0.16))
    echelle = 11.9 / sum(largeurs)
    largeurs = [l * echelle for l in largeurs]
    haut = HAUT
    if bloc.get("intro"):
        # la phrase d'introduction au-dessus du tableau, qui descend d'autant
        boite = slide.shapes.add_textbox(Inches(GAUCHE + 0.15), Inches(1.35), Inches(11.6), Inches(0.8))
        boite.name = "Introduction"
        boite.text_frame.word_wrap = True
        remplir_paragraphe(boite.text_frame.paragraphs[0], " ".join(bloc["intro"]), Pt(22))
        for r in boite.text_frame.paragraphs[0].runs:
            if r.font.color.type is None:
                r.font.color.rgb = BLANC
        haut = 2.2
    haut_ligne = min(0.72, (BAS_PUCES - 0.55 - haut) / len(lignes)) if bloc.get("intro") \
        else min(0.72, (BAS_PUCES - HAUT) / len(lignes))
    taille = 22 if len(lignes) <= 6 else 20
    y = haut
    for i, ligne in enumerate(lignes):
        x = GAUCHE
        if i:
            fond = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                          Inches(11.9), Inches(haut_ligne - 0.08))
            fond.adjustments[0] = 0.18
            fond.fill.solid()
            fond.fill.fore_color.rgb = FOND_CODE
            fond.line.fill.background()
        for c, cellule in enumerate(ligne):
            boite = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y),
                                             Inches(largeurs[c] - 0.3), Inches(haut_ligne - 0.08))
            tf = boite.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            par = tf.paragraphs[0]
            remplir_paragraphe(par, cellule, Pt(taille - (2 if i == 0 else 0)))
            for r in par.runs:
                if i == 0:
                    r.font.bold = True
                    if r.font.color.type is None:
                        r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
                elif r.font.color.type is None:
                    r.font.color.rgb = BLANC
            x += largeurs[c]
        y += haut_ligne
    ajouter_numero(slide, numero)
    return slide


RENVOI = re.compile(r"^\*\*(?P<code>[A-Z]+\.\d+[a-z]?)\*\* - (?P<titre>.*?)"
                    r"(?:\s*\*\((?P<src>[^)]*)\)\*)?$")


def slide_renvoi(prs, bloc, numero, image_exo, fascicules):
    slide = prs.slides.add_slide(layout(prs, "Titre seul"))
    titre(slide, bloc["titre"])
    cartes, extras = [], []
    for p in bloc["puces"]:
        m = RENVOI.match(p)
        (cartes if m else extras).append(m or p)
    y = 1.55
    for m in cartes:
        code = m.group("code")
        source = m.group("src") or fascicules.get(code.split(".")[0], "")
        carte = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(GAUCHE), Inches(y),
                                       Inches(10.2), Inches(1.05))
        carte.adjustments[0] = 0.12
        carte.fill.solid()
        carte.fill.fore_color.rgb = FOND_CODE
        carte.line.fill.background()
        etiquette = slide.shapes.add_textbox(Inches(0.95), Inches(y), Inches(2.1), Inches(1.05))
        tf = etiquette.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run()
        r.text = code
        r.font.name = "Consolas"
        r.font.bold = True
        r.font.size = Pt(30)
        r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
        texte = slide.shapes.add_textbox(Inches(3.15), Inches(y), Inches(7.6), Inches(1.05))
        tf = texte.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf.paragraphs[0]
        remplir_paragraphe(p1, m.group("titre"), Pt(22))
        for run in p1.runs:
            run.font.color.rgb = BLANC
        if source:
            p2 = tf.add_paragraph()
            r = p2.add_run()
            r.text = source
            r.font.size = Pt(14)
            r.font.color.rgb = COULEURS["texte"]
        y += 1.25
    if extras:
        boite = slide.shapes.add_textbox(Inches(GAUCHE), Inches(y + 0.1), Inches(10.2), Inches(0.9))
        tf = boite.text_frame
        tf.word_wrap = True
        for i, e in enumerate(extras):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            remplir_paragraphe(par, e, Pt(18))
    if image_exo:
        slide.shapes.add_picture(io.BytesIO(image_exo["blob"]), Inches(11.3), Inches(0.2),
                                 Inches(1.4), Inches(1.4 * image_exo["h"] / image_exo["l"]))
    # icones a cote de leurs cartes, un groupe par sorte, animation s'il y en a deux
    # (meme logique que _outils/organiser_exercices.py, qui la porte)
    organiser_exercices(slide, LOGO_GAMEHUB if LOGO_GAMEHUB.exists() else None)
    ajouter_numero(slide, numero)
    return slide


def slide_demo(prs, bloc, numero, icone, document):
    slide = prs.slides.add_slide(layout(prs, "Titre seul"))
    titre(slide, "Démonstration")
    sujet = puces_visibles(bloc)[0].replace("Démo : ", "")
    cadre = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(2.0),
                                   Inches(10.1), Inches(3.6))
    cadre.adjustments[0] = 0.08
    cadre.fill.solid()
    cadre.fill.fore_color.rgb = FOND_CODE
    cadre.line.fill.background()
    if icone:
        slide.shapes.add_picture(io.BytesIO(icone), Inches(2.1), Inches(2.75),
                                 Inches(2.1), Inches(2.1))
    boite = slide.shapes.add_textbox(Inches(4.6), Inches(2.3), Inches(6.8), Inches(3.0))
    tf = boite.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Démonstration"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = sujet[0].upper() + sujet[1:]
    r.font.size = Pt(30)
    r.font.color.rgb = BLANC
    p = tf.add_paragraph()
    p.space_before = Pt(12)
    r = p.add_run()
    r.text = f"Déroulé : {document}"
    r.font.size = Pt(14)
    r.font.color.rgb = COULEURS["texte"]
    ajouter_numero(slide, numero)
    return slide


def slide_titre(prs, gabarit, bloc):
    slide = importer_slide(prs, gabarit)
    boites = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    boites.sort(key=lambda s: s.top)
    sous = boites[-1]
    tf = sous.text_frame
    puces = puces_visibles(bloc)
    poser_une_ligne(sous, puces[0])
    for texte in puces[1:]:
        p = tf.add_paragraph()
        p.text = texte
    return slide


# ---------------------------------------------------------------------------
#  Ressources prelevees dans les decks sources
# ---------------------------------------------------------------------------

def image_de(slide, nom):
    for f in slide.shapes:
        if f.name == nom and f.shape_type == 13:
            return {"blob": f.image.blob, "l": f.width, "h": f.height}
    return None


# ---------------------------------------------------------------------------
#  Programme principal
# ---------------------------------------------------------------------------

def construire(cfg, sortie):
    blocs = lire_scenario(cfg.SCENARIO)
    prs = Presentation(str(cfg.GABARIT))
    sources = {nom: Presentation(str(c)) for nom, c in cfg.SOURCES.items()}
    sources["gabarit"] = prs

    def slide_source(ref):
        deck, num = ref
        return sources[deck].slides[num - 1]

    modele_numero = None
    for sl in prs.slides:
        for ph in sl.placeholders:
            if ph.placeholder_format.type is not None \
                    and "SLIDE_NUMBER" in str(ph.placeholder_format.type):
                modele_numero = ph._element
                break
        if modele_numero is not None:
            break

    icone = image_de(slide_source(cfg.ICONE_DEMO[:2]), cfg.ICONE_DEMO[2])
    icone = icone["blob"] if icone else None
    image_exo = image_de(slide_source(cfg.IMAGE_EXO[:2]), cfg.IMAGE_EXO[2])

    ordre, rapport = [], []
    for ident, action, arg in cfg.PLAN:
        bloc = blocs[ident]
        if action == "titre":
            slide = slide_titre(prs, slide_source(cfg.GABARIT_TITRE), bloc)
        elif action == "inter":
            slide = importer_slide(prs, slide_source(cfg.GABARIT_INTERCALAIRE))
            image = getattr(cfg, "IMAGES_INTERCALAIRE", {}).get(ident)
            if image:
                remplacer_fond(slide, image)
            textes_intercalaire(slide, *cfg.INTERCALAIRES[ident])
        elif action == "import":
            slide = importer_slide(prs, slide_source(arg))
        elif action == "new":
            slide = slide_contenu(prs, bloc, modele_numero)
        elif action == "table":
            slide = slide_tableau(prs, bloc, modele_numero)
        elif action == "demo":
            slide = slide_demo(prs, bloc, modele_numero, icone, arg)
        elif action == "renvoi":
            slide = slide_renvoi(prs, bloc, modele_numero, image_exo, cfg.FASCICULES)
        else:
            raise ValueError(f"action inconnue : {action}")
        if hasattr(cfg, "retoucher"):
            cfg.retoucher(slide, ident, bloc)
        ecrire_notes(slide, bloc["note"])
        ordre.append(slide)
        rapport.append((ident, action, len(bloc["puces"]), len(bloc["note"])))

    lst = prs.slides._sldIdLst
    ids = [s.slide_id for s in ordre]
    elements = {int(el.get("id")): el for el in lst}
    for sid, el in elements.items():
        lst.remove(el)
        if sid not in ids:
            prs.part.drop_rel(el.get(qn("r:id")))
    for sid in ids:
        lst.append(elements[sid])

    prs.save(str(sortie))
    return rapport


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    cfg = importlib.import_module(sys.argv[1])
    if not hasattr(cfg, "PLAN"):
        sys.exit(f"{sys.argv[1]} n'a pas de PLAN : ce module ne se construit pas avec "
                 f"construire_deck.py"
                 + (f" mais avec {cfg.MOTEUR}" if hasattr(cfg, "MOTEUR") else "") + ".")
    sortie = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(f"{cfg.CODE}_deck.pptx")
    rapport = construire(cfg, sortie)
    for ligne in rapport:
        print(ligne)
    print(f"{len(rapport)} slides -> {sortie}")


if __name__ == "__main__":
    main()
