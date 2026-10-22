#!/usr/bin/env python3
"""Finalise le deck d'un module : sommaire, sections numerotees, reperes.

Commun a tous les modules ; ce qui differe de l'un a l'autre vit dans
le fichier module_xxx.py de chaque module. Travaille sur le deck tel qu'il
sort de PowerPoint, retouches comprises : il ne reconstruit rien, il ajoute
ou remplace. Il peut etre relance sans risque : les numeros et les reperes
existants sont d'abord retires, puis reposes ; le sommaire existant est
vide puis redessine.

Usage :
    python3 finaliser_deck.py MODULE entree.pptx sortie.pptx [dossier_captures]

MODULE est le code d'un module (TMP, ITF...) : sa configuration est lue dans
module_xxx.py, a cote de ce script. Un code inconnu affiche la liste des
modules disponibles.

Le dossier de captures ne sert qu'a POO. Il doit contenir
A__Javadoc_Scanner.png, B__Javadoc_nextInt.png et C__Javadoc_nextLine.png.
Sans lui, les captures ne sont pas touchees.
"""
import copy
import importlib.util
import re
import sys
import uuid
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mise_en_forme import cran_dessous, uniformiser_titre_intercalaire  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "_outils"))
from animer_code import animer_slide as animer_code  # noqa: E402
from animer_sommaire import animer_slide as animer_sommaire  # noqa: E402
from animer_lignes import animer_slide as animer_lignes  # noqa: E402
from ajuster_code import ajuster_slide as ajuster_code  # noqa: E402
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
#  Configuration par module
# ---------------------------------------------------------------------------
# Ce qui est propre a un module vit dans son fichier module_xxx.py, a cote de
# ce script. Le finaliseur y lit :
#   CODE      le code du module, ex. "TMP"
#   ORDRE     le rang du deck dans le cours (1 = INI, 2 = POO, 3 = TMP...)
#   SECTIONS  le texte des intercalaires, numero retire
#   LIBELLES  la version courte, pour le repere en bas de slide et le sommaire
# Ajouter un module ne demande donc jamais de modifier ce fichier.

ICI = Path(__file__).resolve().parent
CODE_VALIDE = re.compile(r"^[A-Za-z0-9]+$")


class ConfigInvalide(Exception):
    pass


def charger_config(chemin):
    """Charge un module_xxx.py par son chemin, sans dependre du dossier courant."""
    spec = importlib.util.spec_from_file_location(chemin.stem, chemin)
    cfg = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(cfg)
    except Exception as e:
        raise ConfigInvalide(f"{chemin.name} ne se charge pas : {type(e).__name__} : {e}")
    return cfg


def verifier_config(cfg, code, nom):
    """Leve ConfigInvalide si les donnees du finaliseur manquent ou se contredisent."""
    for attr in ("SECTIONS", "LIBELLES"):
        valeur = getattr(cfg, attr, None)
        if not isinstance(valeur, (list, tuple)) or not valeur \
                or not all(isinstance(x, str) and x.strip() for x in valeur):
            raise ConfigInvalide(f"{nom} : {attr} manque, est vide ou contient "
                                 f"autre chose que du texte")
    if len(cfg.SECTIONS) != len(cfg.LIBELLES):
        raise ConfigInvalide(f"{nom} : {len(cfg.SECTIONS)} SECTIONS pour "
                             f"{len(cfg.LIBELLES)} LIBELLES")
    if getattr(cfg, "CODE", code) != code:
        raise ConfigInvalide(f"{nom} declare CODE = {cfg.CODE!r}, attendu {code!r}")
    if not isinstance(getattr(cfg, "ORDRE", None), int):
        raise ConfigInvalide(f"{nom} : ORDRE manque (le rang du deck dans le cours)")


def modules_connus():
    """[(ordre, code)] des module_*.py valides, dans l'ordre du cours.

    L'ordre vient d'ORDRE, jamais de l'ordre de lecture du disque. Les
    fichiers invalides sont signales et ecartes.
    """
    connus, problemes = [], []
    for chemin in sorted(ICI.glob("module_*.py")):
        code = chemin.stem[len("module_"):].upper()
        try:
            cfg = charger_config(chemin)
            verifier_config(cfg, code, chemin.name)
        except ConfigInvalide as e:
            problemes.append(str(e))
            continue
        connus.append((cfg.ORDRE, code))
    connus.sort()
    return connus, problemes


def lister_modules():
    connus, problemes = modules_connus()
    lignes = [f"  {o:>3}  {c}" for o, c in connus]
    rangs = [o for o, _ in connus]
    doublons = sorted({o for o in rangs if rangs.count(o) > 1})
    for o in doublons:
        lignes.append(f"  attention : ORDRE {o} est pris par "
                      + ", ".join(c for r, c in connus if r == o))
    for p in problemes:
        lignes.append(f"  ecarte : {p}")
    return "\n".join(lignes)


MODULE = SECTIONS = LIBELLES = None


def configurer(code):
    global MODULE, SECTIONS, LIBELLES
    if not CODE_VALIDE.match(code):
        sys.exit(f"Code de module invalide : {code!r}\nModules connus :\n{lister_modules()}")
    chemin = ICI / f"module_{code.lower()}.py"
    if not chemin.exists():
        sys.exit(f"Module inconnu : {code} (aucun fichier {chemin.name} a cote de "
                 f"{Path(__file__).name})\nModules connus :\n{lister_modules()}")
    try:
        cfg = charger_config(chemin)
        verifier_config(cfg, code, chemin.name)
    except ConfigInvalide as e:
        sys.exit(f"Configuration invalide : {e}")
    MODULE = code
    SECTIONS = list(cfg.SECTIONS)
    LIBELLES = list(cfg.LIBELLES)
    avertir_intercalaires(cfg, chemin.name)


def avertir_intercalaires(cfg, nom):
    """Signale un ecart entre INTERCALAIRES (construction) et SECTIONS (finalisation).

    Un intercalaire dont le titre n'est pas dans SECTIONS ne serait ni
    numerote ni repere : c'est l'erreur que le commentaire de chaque
    module_xxx.py demandait d'eviter a la main.
    """
    inter = getattr(cfg, "INTERCALAIRES", None)
    if not inter:
        return
    titres = [inter[k][0] for k in sorted(inter, key=lambda k: [int(x) for x in k.split(".")])]
    if [t.lower() for t in titres] != [s.lower() for s in cfg.SECTIONS]:
        print(f"Attention : dans {nom}, les titres d'INTERCALAIRES ne correspondent "
              f"pas a SECTIONS.\n  INTERCALAIRES : {titres}\n  SECTIONS      : "
              f"{list(cfg.SECTIONS)}", file=sys.stderr)


NOMBRES = {1: "une", 2: "deux", 3: "trois", 4: "quatre", 5: "cinq", 6: "six",
           7: "sept", 8: "huit", 9: "neuf", 10: "dix"}

TITRE_SOMMAIRE = "Au programme"
NOM_REPERE = "Repere de section"
NOM_SOMMAIRE = "Sommaire du module"
NOM_SOMMAIRE_TEXTE = "Sommaire du module, texte"
PREFIXE = re.compile(r"^\d+\.\s+")

FOND = RGBColor(0x1D, 0x1F, 0x22)
GRIS = RGBColor(0xBC, 0xBE, 0xC4)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------------------
#  Reperage
# ---------------------------------------------------------------------------

def titre_intercalaire(slide):
    """Zone de titre d'un intercalaire (image pleine page + zone de texte), ou None."""
    images = [s for s in slide.shapes if s.shape_type == 13
              and s.width > Inches(12) and s.height > Inches(7)]
    if not images:
        return None
    boites = [s for s in slide.shapes if s.has_text_frame and not s.is_placeholder
              and s.text_frame.text.strip()]
    if not boites:
        return None
    return min(boites, key=lambda s: (s.top, s.left))


def numero_section(texte):
    nu = PREFIXE.sub("", texte.strip())
    for i, nom in enumerate(SECTIONS, 1):
        if nu.lower() == nom.lower():
            return i
    return None


def est_sommaire(slide):
    t = slide.shapes.title
    return t is not None and t.text == TITRE_SOMMAIRE


def placeholder_numero(prs):
    for slide in prs.slides:
        for ph in slide.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                return ph._element
    return None


# ---------------------------------------------------------------------------
#  Nettoyage (pour pouvoir relancer)
# ---------------------------------------------------------------------------

def nettoyer(prs):
    lst = prs.slides._sldIdLst
    for slide in prs.slides:
        if est_sommaire(slide):
            retirer(slide, lambda x: x.name in (NOM_SOMMAIRE, NOM_SOMMAIRE_TEXTE))
    for slide in prs.slides:
        for s in list(slide.shapes):
            if s.name == NOM_REPERE:
                s._element.getparent().remove(s._element)
        boite = titre_intercalaire(slide)
        if boite is not None:
            for p in boite.text_frame.paragraphs:
                if p.runs:
                    p.runs[0].text = PREFIXE.sub("", p.runs[0].text.lstrip())
                    break


# ---------------------------------------------------------------------------
#  Captures Javadoc
# ---------------------------------------------------------------------------

def image(slide, chemin, x, y, hauteur_max, largeur_max):
    l, h = Image.open(chemin).size
    ratio = min(largeur_max / (l / 96), hauteur_max / (h / 96))
    larg, haut = l / 96 * ratio, h / 96 * ratio
    pic = slide.shapes.add_picture(str(chemin), Inches(x), Inches(y), Inches(larg), Inches(haut))
    pic.name = f"Capture {Path(chemin).stem}"
    return pic


def retirer(slide, predicat):
    for s in list(slide.shapes):
        if predicat(s):
            s._element.getparent().remove(s._element)


def ecrire_notes(slide, paragraphes):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    for i, texte in enumerate(paragraphes):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for morceau in re.split(r"(`[^`]+`)", texte):
            if not morceau:
                continue
            r = par.add_run()
            if morceau.startswith("`"):
                r.text = morceau[1:-1]
                r.font.name = "Consolas"
            else:
                r.text = morceau


def puces(slide, x, y, largeur, hauteur, lignes, taille=20):
    boite = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(largeur), Inches(hauteur))
    tf = boite.text_frame
    tf.word_wrap = True
    for i, texte in enumerate(lignes):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(10)
        for morceau in re.split(r"(`[^`]+`)", texte):
            if not morceau:
                continue
            r = par.add_run()
            r.font.size = Pt(taille)
            if morceau.startswith("`"):
                r.text = morceau[1:-1]
                r.font.name = "Consolas"
                r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
            else:
                r.text = morceau
                r.font.color.rgb = BLANC
    return boite


def poser_captures(prs, dossier):
    dossier = Path(dossier)
    a = dossier / "A__Javadoc_Scanner.png"
    b = dossier / "B__Javadoc_nextInt.png"
    c = dossier / "C__Javadoc_nextLine.png"
    cible_a = cible_b = None
    for slide in prs.slides:
        t = slide.shapes.title.text if slide.shapes.title is not None else ""
        if t.startswith("Les méthodes de Scanner") or t == "La documentation de Scanner":
            cible_a = slide
        elif t.startswith("Les principales méthodes de Scanner"):
            cible_b = slide

    if cible_a is not None:
        s = cible_a
        retirer(s, lambda x: x.name.startswith("A SUPPRIMER")
                or x.name.startswith("Capture ")
                or x.name.startswith("Rectangle\xa0: coins arrondis")
                or x.name == "Bulles de la slide"
                or (x.is_placeholder and x.placeholder_format.type == PP_PLACEHOLDER.OBJECT))
        s.shapes.title.text = "La documentation de Scanner"
        image(s, a, 6.35, 1.45, 5.3, 6.5)
        b_ = puces(s, 0.8, 1.6, 5.3, 4.5, [
            "Survol de `Scanner` dans IntelliJ, ou `Ctrl+Q` / `F1` (IDE.14)",
            "Une classe du package `java.util`",
            "Découpe l'entrée en jetons, séparés par des blancs",
            "Une méthode `next...` par type à lire",
        ])
        b_.name = "Bulles de la slide"
        ecrire_notes(s, [
            "IntelliJ affiche la documentation d'une classe au simple survol de son nom. "
            "Pour `Scanner`, la bulle indique le package, `java.util`, puis le principe : "
            "l'entrée est découpée en jetons, séparés par défaut par des espaces ou des "
            "passages à la ligne, et chaque jeton peut être converti dans le type voulu "
            "grâce aux différentes méthodes `next`.",
            "Les exemples de la bulle lisent la console et un fichier. Pour ce cours, "
            "quelques méthodes suffisent : `nextLine()` lit une ligne entière, `next()` "
            "lit un mot, `nextInt()` et `nextDouble()` lisent un nombre.",
        ])

    if cible_b is not None:
        s = cible_b
        retirer(s, lambda x: x.name.startswith("A SUPPRIMER")
                or x.name.startswith("Capture ")
                or (x.shape_type == 13 and not x.name.startswith("Capture ")))
        corps = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1]
        if corps:
            corps[0].top, corps[0].height = Inches(1.5), Inches(1.3)
        image(s, b, 1.3, 2.85, 3.85, 5.3)
        image(s, c, 6.9, 2.85, 3.85, 5.3)
        ecrire_notes(s, [
            "La bulle d'une méthode donne sa signature, ce qu'elle retourne et les "
            "exceptions qu'elle peut lever.",
            "`nextInt()` lit le jeton suivant et le retourne sous forme d'entier. Elle lève "
            "une `InputMismatchException` si ce jeton n'est pas un entier : c'est ce qui "
            "arrive quand l'utilisateur tape des lettres là où le programme attend un "
            "nombre. La bulle signale aussi la variante `nextInt(radix)`, qui interprète "
            "le nombre dans la base indiquée.",
            "`nextLine()` retourne le reste de la ligne courante, sans le passage à la ligne "
            "final, et se place au début de la ligne suivante. C'est ce comportement qui "
            "explique le piège de la slide suivante.",
        ])


# ---------------------------------------------------------------------------
#  Sommaire et numerotation
# ---------------------------------------------------------------------------

def carte(slide, x, y, l, h):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(l), Inches(h))
    c.adjustments[0] = 0.15
    c.fill.solid()
    c.fill.fore_color.rgb = FOND
    c.line.fill.background()
    return c


def ajouter_sommaire(prs, numero):
    existante = [sl for sl in prs.slides if est_sommaire(sl)]
    if existante:
        slide = existante[0]
    else:
        layout = [l for l in prs.slide_layouts if l.name == "Titre seul"][0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = TITRE_SOMMAIRE
    une_colonne = len(LIBELLES) <= 4
    largeur = 12.0 if une_colonne else 5.9
    for i, libelle in enumerate(LIBELLES):
        col, rang = (0, i) if i < 4 else (1, i - 4)
        x, y = 0.7 + col * 6.1, 1.6 + rang * 1.2
        c = carte(slide, x, y, largeur, 1.0)
        c.name = NOM_SOMMAIRE
        boite = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y), Inches(largeur - 0.4), Inches(1.0))
        boite.name = NOM_SOMMAIRE_TEXTE
        tf = boite.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = f"{i + 1}   "
        r.font.size = Pt(30); r.font.bold = True
        r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
        r = p.add_run(); r.text = libelle
        r.font.size = Pt(26); r.font.color.rgb = BLANC
    if numero is not None and not existante:
        slide.shapes._spTree.append(copy.deepcopy(numero))
    ecrire_notes(slide, [
        f"Le module compte {NOMBRES.get(len(LIBELLES), len(LIBELLES))} sections : " + " ; ".join(
            f"{i + 1}. {l}" for i, l in enumerate(LIBELLES)) + ".",
        "Chaque intercalaire porte le numéro de sa section, et chaque page rappelle en "
        "bas à droite la section en cours.",
    ])
    if not existante:
        # le sommaire vient juste apres la page de titre
        lst = prs.slides._sldIdLst
        el = lst[-1]
        lst.remove(el)
        lst.insert(1, el)
    return slide


def numeroter(prs):
    courante = None
    for slide in prs.slides:
        boite = titre_intercalaire(slide)
        if boite is not None:
            n = numero_section(boite.text_frame.text)
            if n is not None:
                courante = n
                run = next(p.runs[0] for p in boite.text_frame.paragraphs if p.runs)
                run.text = f"{n}. " + run.text.lstrip()
                elargir(boite, run)
                continue
        if courante is None or est_sommaire(slide):
            continue
        if encombre(slide, Inches(7.2), Inches(6.98), Inches(4.9), Inches(0.35)):
            continue
        repere = slide.shapes.add_textbox(Inches(7.2), Inches(6.98), Inches(4.9), Inches(0.35))
        repere.name = NOM_REPERE
        tf = repere.text_frame
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = f"{MODULE} · "
        r.font.size = Pt(11); r.font.color.rgb = GRIS
        r = p.add_run(); r.text = f"{courante}. {LIBELLES[courante - 1]}"
        r.font.size = Pt(11); r.font.bold = True
        r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4


def encombre(slide, x, y, l, h):
    """Vrai si une forme visible occupe deja l'emplacement du repere."""
    for f in slide.shapes:
        if f.is_placeholder and f.placeholder_format.type in (
                PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE):
            continue
        if f.left is None or f.left >= Inches(13.333) or f.top >= Inches(7.5):
            continue   # formes rangees hors de la slide (bulles animees)
        if f.left < x + l and f.left + f.width > x and f.top < y + h and f.top + f.height > y:
            return True
    return False


# ---------------------------------------------------------------------------
#  Taille du texte : 28 pt quand il tient, moins quand il est trop long
# ---------------------------------------------------------------------------

TAILLE_CIBLE = 28
TAILLES = (28, 26, 24, 22, 20)

# Consolas est plus large que Corbel : dans du texte courant, un mot de code
# se pose quatre points plus bas que la prose qui l'entoure (28 -> 24).


def taille_inline(t):
    """Taille d'un mot de code dans de la prose de taille t : un cran en dessous."""
    return cran_dessous(t)


def est_code(run):
    return (run.font.name or "") == "Consolas"
HAUT, BAS = 1.5, 6.75


def hauteur_texte(paragraphes, largeur_po, taille, espace_pt=10):
    car_par_ligne = max(1, int((largeur_po - 0.45) * 72 / (0.55 * taille)))
    lignes = sum(max(1, -(-len(t) // car_par_ligne)) for t in paragraphes)
    return lignes * taille * 1.2 / 72 + len(paragraphes) * espace_pt / 72 + 0.15


def tailles_runs(forme):
    return {r.font.size.pt for p in forme.text_frame.paragraphs for r in p.runs if r.font.size}


def poser_taille(forme, taille, seulement=None, code_a_part=True):
    for p in forme.text_frame.paragraphs:
        for r in p.runs:
            code = code_a_part and est_code(r)
            if seulement is not None:
                attendu = {taille_inline(x) for x in seulement} if code else set(seulement)
                if not (r.font.size and r.font.size.pt in attendu):
                    continue
            r.font.size = Pt(taille_inline(taille) if code else taille)


def harmoniser_code_inline(prs):
    """Aligne chaque mot de code sur la prose qui l'entoure, partout.

    Une zone nommee Code est entierement du code : on n'y touche pas. Un
    paragraphe entierement en Consolas non plus - c'est un affichage voulu,
    comme l'etiquette d'une carte d'exercice.
    """
    n = 0
    for slide in prs.slides:
        for forme in slide.shapes:
            if not forme.has_text_frame or forme.name == "Code":
                continue
            for par in forme.text_frame.paragraphs:
                prose = [r for r in par.runs if not est_code(r) and r.font.size]
                codes = [r for r in par.runs if est_code(r)]
                if not prose or not codes:
                    continue
                cible = Pt(taille_inline(max(r.font.size.pt for r in prose)))
                for r in codes:
                    if r.font.size != cible:
                        r.font.size = cible
                        n += 1
    return n


def ajuster_contenu(slide):
    """Slides de contenu construites par le script : puces en 20 ou 22 pt."""
    corps = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1]
    if not corps or not tailles_runs(corps[0]) or not tailles_runs(corps[0]) <= {20.0, 22.0}:
        return False
    corps = corps[0]
    textes = [p.text for p in corps.text_frame.paragraphs if p.text.strip()]
    code = [f for f in slide.shapes if f.name == "Code"]
    code = code[0] if code else None
    a_droite = code is not None and code.left > Inches(3)
    for t in TAILLES:
        if code is None:
            if hauteur_texte(textes, Emu(corps.width).inches, t) <= BAS - HAUT:
                break
        elif a_droite:
            largeur = Emu(code.left).inches - 0.25 - Emu(corps.left).inches
            if hauteur_texte(textes, largeur, t) <= BAS - HAUT:
                break
        else:
            # code sous les puces : on accepte de reduire le code jusqu'a 16 pt
            h = hauteur_texte(textes, Emu(corps.width).inches, t)
            dispo = BAS - HAUT - h - 0.2
            if taille_code_pour(code, dispo) >= 16:
                break
    poser_taille(corps, t)
    if code is not None and not a_droite:
        h = hauteur_texte(textes, Emu(corps.width).inches, t)
        corps.height = Inches(h)
        dispo = BAS - HAUT - h - 0.2
        tc = taille_code_pour(code, dispo)
        n = len(code.text_frame.paragraphs)
        poser_taille(code, tc, code_a_part=False)
        code.top = Inches(HAUT + h + 0.2)
        code.height = Inches(min(dispo, n * tc * 1.2 / 72 + 0.3))
    return True


def taille_code_pour(code, hauteur_po):
    """Plus grande taille de code (sans depasser l'actuelle) qui tient en hauteur."""
    actuelle = max(tailles_runs(code) or {18.0})
    n = len(code.text_frame.paragraphs)
    possible = (hauteur_po - 0.3) * 72 / (1.2 * n)
    return min(actuelle, int(possible * 2) / 2)


def ajuster_tailles(prs):
    for slide in prs.slides:
        titre = slide.shapes.title.text if slide.shapes.title is not None else ""
        if ajuster_contenu(slide):
            continue
        boites = [f for f in slide.shapes if f.has_text_frame and not f.is_placeholder]
        if titre.startswith("Exercice"):
            for b in boites:
                tailles = tailles_runs(b)
                if 22.0 in tailles:
                    longueur = len(b.text_frame.paragraphs[0].text)
                    poser_taille(b, 28 if longueur <= 34 else 24, {22.0})
                elif tailles == {18.0}:
                    poser_taille(b, 24, {18.0})
        elif titre == "Les quatre principes de la POO":
            for b in boites:
                poser_taille(b, 24, {20.0})
                poser_taille(b, 18, {16.0})
        elif titre == "Du code source au programme":
            for b in boites:
                poser_taille(b, 28, {20.0})
        elif titre == "La documentation de Scanner":
            for b in boites:
                if b.name == "Bulles de la slide" and tailles_runs(b) == {20.0}:
                    textes = [p.text for p in b.text_frame.paragraphs]
                    t = next(t for t in TAILLES
                             if hauteur_texte(textes, Emu(b.width).inches, t) <= BAS - HAUT)
                    poser_taille(b, t)


# ---------------------------------------------------------------------------
#  Sections PowerPoint (volet de navigation)
# ---------------------------------------------------------------------------
# PowerPoint exige que chaque slide appartienne a une section des qu'il en
# existe : une slide orpheline declenche la demande de reparation.

URI_SECTIONS = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def reconstruire_sections(prs):
    groupes = [["Ouverture", []]]
    for slide in prs.slides:
        boite = titre_intercalaire(slide)
        n = numero_section(boite.text_frame.text) if boite is not None else None
        if n is not None:
            groupes.append([f"{n}. {LIBELLES[n - 1]}", []])
        groupes[-1][1].append(slide.slide_id)

    pres = prs.part._element
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    for ext in ext_lst.findall(qn("p:ext")):
        if ext.get("uri") == URI_SECTIONS:
            ext_lst.remove(ext)
    ext = etree.Element(qn("p:ext"), uri=URI_SECTIONS)
    ext_lst.insert(0, ext)
    lst = etree.SubElement(ext, f"{{{NS_P14}}}sectionLst", nsmap={"p14": NS_P14})
    for nom, ids in groupes:
        guid = "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, f"{MODULE}/{nom}")).upper() + "}"
        sec = etree.SubElement(lst, f"{{{NS_P14}}}section", name=nom, id=guid)
        sld = etree.SubElement(sec, f"{{{NS_P14}}}sldIdLst")
        for i in ids:
            etree.SubElement(sld, f"{{{NS_P14}}}sldId", id=str(i))


# ---------------------------------------------------------------------------
#  Titres de meme taille, zones de code a coins arrondis
# ---------------------------------------------------------------------------

TAILLE_TITRE = 40
FOND_CODE = "1D1F22"
RAYON_CODE = 0.1   # pouces


def uniformiser_titres(prs):
    for slide in prs.slides:
        t = slide.shapes.title
        if t is None or t.placeholder_format.type != PP_PLACEHOLDER.TITLE:
            continue
        body_pr = t.text_frame._txBody.find(qn("a:bodyPr"))
        for enfant in list(body_pr):
            if enfant.tag in (qn("a:normAutofit"), qn("a:spAutoFit"), qn("a:noAutofit")):
                body_pr.remove(enfant)
        etree.SubElement(body_pr, qn("a:noAutofit"))
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(TAILLE_TITRE)


def uniformiser_intercalaires(prs):
    """Meme style de titre pour tous les intercalaires (voir mise_en_forme.py)."""
    tailles = []
    for slide in prs.slides:
        boite = titre_intercalaire(slide)
        if boite is not None:
            tailles.append(uniformiser_titre_intercalaire(boite, prs.slide_width))
    return tailles


def arrondir_code(prs):
    for slide in prs.slides:
        for f in slide.shapes:
            sppr = f._element.find(qn("p:spPr"))
            if sppr is None:
                continue
            couleur = sppr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            geom = sppr.find(qn("a:prstGeom"))
            if couleur is None or geom is None or couleur.get("val", "").upper() != FOND_CODE:
                continue
            if geom.get("prst") not in ("rect", "roundRect"):
                continue
            cote = min(f.width, f.height)
            adj = min(50000, int(Inches(RAYON_CODE) * 100000 / cote)) if cote else 0
            geom.set("prst", "roundRect")
            av = geom.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(geom, qn("a:avLst"))
            for g in list(av):
                av.remove(g)
            etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {adj}")


# L'ombre des boites noires : petite, portee vers le bas a droite, noire a
# 40 %. C'est celle posee a la main sur le deck INI ; tous les decks la
# recoivent, y compris les formes ajoutees plus tard dans PowerPoint.
OMBRE_BOITE = (
    '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:outerShdw blurRad="50800" dist="38100" dir="2700000" algn="tl" rotWithShape="0">'
    '<a:prstClr val="black"><a:alpha val="40000"/></a:prstClr>'
    '</a:outerShdw></a:effectLst>'
)


def ombrer_boites(prs):
    """Pose OMBRE_BOITE sur chaque boite a fond 1D1F22, groupes compris."""
    modele = etree.fromstring(OMBRE_BOITE)
    for slide in prs.slides:
        for sp in slide._element.iter(qn("p:sp")):
            sppr = sp.find(qn("p:spPr"))
            if sppr is None:
                continue
            couleur = sppr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            if couleur is None or couleur.get("val", "").upper() != FOND_CODE:
                continue
            for ancien in sppr.findall(qn("a:effectLst")):
                sppr.remove(ancien)
            nouveau = copy.deepcopy(modele)
            # ordre du schema : xfrm, geometrie, remplissage, ln, effectLst...
            ln = sppr.find(qn("a:ln"))
            if ln is not None:
                ln.addnext(nouveau)
            else:
                sppr.append(nouveau)


def nettoyer_animations(prs):
    """Retire les animations qui visent une forme disparue.

    Une animation dont la cible n'existe plus declenche, elle aussi, la
    demande de reparation de PowerPoint.
    """
    for i, slide in enumerate(prs.slides, 1):
        timing = slide._element.find(qn("p:timing"))
        if timing is None:
            continue
        ids = {str(f.shape_id) for f in slide.shapes}
        for f in slide.shapes:
            if f.shape_type == 6:   # groupe : ses enfants peuvent etre animes
                ids |= {el.get("id") for el in f._element.iter(qn("p:cNvPr"))}
        cibles = {el.get("spid") for el in timing.iter(qn("p:spTgt"))}
        if cibles - ids:
            slide._element.remove(timing)
            print(f"  slide {i} : animations retirees (cibles disparues : "
                  f"{', '.join(sorted(cibles - ids))})")


def elargir(boite, run):
    taille = run.font.size.pt if run.font.size else 60
    texte = boite.text_frame.text
    voulu = Inches(len(texte) * taille * 0.5 / 72 + 0.6)
    maxi = Inches(13.333 - 0.3) - boite.left
    largeur = min(voulu, maxi)
    if largeur > boite.width:
        boite.width = Emu(int(largeur))


def main():
    if len(sys.argv) < 4:
        sys.exit(f"{__doc__}\nModules connus :\n{lister_modules()}")
    configurer(sys.argv[1].upper())
    entree, sortie = sys.argv[2], sys.argv[3]
    captures = sys.argv[4] if len(sys.argv) > 4 else None
    prs = Presentation(entree)
    nettoyer(prs)
    if captures and MODULE == "POO":
        poser_captures(prs, captures)
    ajouter_sommaire(prs, placeholder_numero(prs))
    ajuster_tailles(prs)
    uniformiser_titres(prs)
    arrondir_code(prs)
    ombrer_boites(prs)
    numeroter(prs)
    reconstruire_sections(prs)
    uniformiser_intercalaires(prs)    # apres numeroter : le numero compte dans la largeur
    nettoyer_animations(prs)
    n = harmoniser_code_inline(prs)
    for slide in prs.slides:          # le code sous le texte vient au clic, en fondu
        ajuster_code(slide, prs.slide_width)   # code trop large : retreci et centre
        animer_code(slide)
        animer_sommaire(slide)        # « Au programme » : une section par clic
        animer_lignes(slide)          # tableau en bandeaux : une ligne par clic
    print(f"{n} mots de code realignes sur leur prose")
    prs.save(sortie)
    print(f"{len(prs.slides)} slides -> {sortie}")


if __name__ == "__main__":
    main()
