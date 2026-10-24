#!/usr/bin/env python3
"""Configuration du module REC pour construire_deck.py et finaliser_deck.py.

Arborescence attendue, relative au dossier de ce fichier :
    ../gabarit_deck.pptx       gabarit commun : theme, dispositions, 4 slides modeles
    ../sources/d02.pptx        ancien deck 02, qui fournit l'intercalaire pleine page
    ../scenario/Scenario_REC.md   le scenario

Un deck court (24/10/2026) : il presente les deux series d'exercices
recapitulatifs, a faire chez soi pendant les semaines sans cours. Aucune
notion nouvelle : la matiere est celle des fascicules REC.

Usage :
    python3 construire_deck.py module_rec REC_deck.pptx
    python3 finaliser_deck.py REC REC_deck.pptx REC_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent.parent   # _slides

CODE = "REC"
ORDRE = 16   # rang du deck dans le cours : apres FIC
SCENARIO = ICI / "scenario" / "Scenario_REC.md"
GABARIT = ICI / "gabarit_deck.pptx"
SOURCES = {"d02": ICI / "sources" / "d02.pptx"}

GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d02", 13)       # « Structure », pleine page, comme PKG et FIC
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"REC": "fascicule REC"}

SECTIONS = [
    "Réviser dans la forme de l'examen",
    "Première série, après JUN",
    "Seconde série, après JavaFX",
]

LIBELLES = [
    "Pourquoi",
    "Série 1",
    "Série 2",
]

INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}
# Les trois images generees le 24/10/2026, dans _illustrations/illus_generees.
IMAGES_INTERCALAIRE = {f"{n}.1": ICI.parent / "_illustrations" / "illus_generees" / f"REC_section{n}.jpg"
                       for n in (1, 2, 3)}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "table", None), ("1.3", "new", None),

    ("2.1", "inter", None),
    ("2.2", "table", None), ("2.3", "renvoi", None), ("2.4", "renvoi", None),

    ("3.1", "inter", None),
    ("3.2", "table", None), ("3.3", "renvoi", None), ("3.4", "renvoi", None),

    ("4.1", "new", None),
]

# Meme intercalaire que PKG et FIC, meme taille reduite : « 1. Réviser dans
# la forme de l'examen » (36 caracteres) passerait a la ligne en 66 pt.
TAILLE_INTERCALAIRE = 48


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
