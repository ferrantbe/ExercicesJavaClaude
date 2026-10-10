#!/usr/bin/env python3
"""Configuration du module GEN pour construire_deck.py et finaliser_deck.py.

Arborescence attendue, relative au dossier de ce fichier :
    sources/gabarit_deck.pptx  gabarit commun : theme, dispositions, 4 slides modeles
    sources/d08.pptx           ancien deck 08 (generiques et lambdas reunis), qui
                               fournit l'intercalaire ; appele d08a dans SOURCES,
                               nom historique conserve
    scenario/Scenario_GEN.md  le scenario

Usage :
    python3 construire_deck.py module_gen GEN_deck.pptx
    python3 finaliser_deck.py GEN GEN_deck.pptx GEN_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent

CODE = "GEN"
ORDRE = 12   # rang du deck dans le cours : apres COL, EXC et JUN
SCENARIO = ICI / "scenario" / "Scenario_GEN.md"
GABARIT = ICI / "sources" / "gabarit_deck.pptx"
SOURCES = {"d08a": ICI / "sources" / "d08.pptx"}

# Slides prelevees dans les decks sources : (deck, numero[, nom de forme])
GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d08a", 8)       # « Les types generiques », meme structure que d06/20
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"GEN": "fascicule GEN", "G": "fil rouge GameHub"}

# Lu par finaliser_deck.py.
# SECTIONS : le texte des intercalaires du deck, numero retire, mot pour mot
#            (casse mise a part) ; un intercalaire absent de cette liste
#            n'est ni numerote ni repere.
# LIBELLES : la version courte, pour le sommaire et le repere en bas de slide.
SECTIONS = [
    "Un type en paramètre",
    "Autant de types qu'il en faut",
    "Ce que le compilateur efface",
    "Une famille de types",
    "Une méthode générique",
    "Hériter d'une classe générique",
]

LIBELLES = [
    "La classe générique",
    "Plusieurs paramètres",
    "L'effacement",
    "Les jokers",
    "Méthodes génériques",
    "L'héritage",
]

# Titre des intercalaires, sans numero : le finaliseur le pose.
# Les titres sont ceux de SECTIONS, dans le meme ordre : ils en sont tires,
# pour qu'un titre change a un seul endroit.
INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}

# Une seule image pour les six sections : celle de l'ancien deck, deja
# dans la slide modele. Frederic les differenciera ensuite dans PowerPoint.
IMAGES_INTERCALAIRE = {}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "new", None),
    ("1.3", "demo", "Demo_GEN"),
    ("1.4", "new", None), ("1.5", "new", None), ("1.6", "renvoi", None),

    ("2.1", "inter", None),
    ("2.2", "new", None), ("2.3", "new", None), ("2.4", "renvoi", None),

    ("3.1", "inter", None),
    ("3.2", "new", None), ("3.3", "new", None), ("3.4", "table", None),
    ("3.5", "new", None),

    ("4.1", "inter", None),
    ("4.2", "new", None), ("4.3", "renvoi", None), ("4.4", "new", None),
    ("4.5", "new", None),

    ("5.1", "inter", None),
    ("5.2", "new", None), ("5.3", "new", None), ("5.4", "renvoi", None),

    ("6.1", "inter", None),
    ("6.2", "new", None), ("6.3", "new", None), ("6.4", "renvoi", None),

    ("7.1", "renvoi", None),
]


# Le plus long titre d'intercalaire, numero compris (« 6. Heriter d'une
# classe generique », 33 caracteres), passerait a la ligne en 60 pt. Meme
# taille reduite que ITF et EXC, pour toutes les sections.
TAILLE_INTERCALAIRE = 44


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
