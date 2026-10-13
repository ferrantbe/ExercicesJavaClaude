"""Regles de mise en forme communes au moteur et aux outils.

Importe par construire_deck.py, finaliser_deck.py et les outils de _outils/.
Une regle ne vit qu'ici : la changer ici la change partout.

- Un mot de code est un cran plus petit que le texte qui l'entoure, sur
  l'echelle de tailles de PowerPoint (40 -> 36, 28 -> 24), valide le 12/10/2026.
- Le titre d'un intercalaire est en Calibri gras, aligne a gauche, a 66 pt ;
  s'il ne tient pas sur une ligne, il descend d'un cran a la fois (60, 54,
  48, 44...) jusqu'a tenir. Valide le 13/10/2026.
"""
from pathlib import Path

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ECHELLE = [8, 9, 10, 10.5, 11, 12, 14, 16, 18, 20, 24, 28, 32, 36, 40, 44,
           48, 54, 60, 66, 72, 80, 88, 96]

TAILLE_INTERCALAIRE = 66
TAILLE_INTERCALAIRE_MIN = 32
POLICE_INTERCALAIRE = "Calibri"
MARGE_DROITE = Inches(0.3)


def cran_dessous(pt):
    """La taille juste en dessous de pt sur l'echelle de PowerPoint."""
    plus_petits = [t for t in ECHELLE if t < pt - 0.01]
    return plus_petits[-1] if plus_petits else pt


# --- mesure du texte --------------------------------------------------------
# Carlito a exactement les metriques de Calibri ; a defaut, on cherche
# Calibri lui-meme, puis on se rabat sur une estimation.
_POLICES_GRAS = [
    "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
    "C:/Windows/Fonts/calibrib.ttf",
    "/Library/Fonts/Calibri Bold.ttf",
    str(Path.home() / "Library/Fonts/Calibri Bold.ttf"),
    "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/calibrib.ttf",
]


def _largeur_calibri_gras(texte, pt):
    """Largeur en pouces du texte en Calibri gras de pt points."""
    try:
        from PIL import ImageFont
        for chemin in _POLICES_GRAS:
            if Path(chemin).exists():
                f = ImageFont.truetype(chemin, 1000)
                return f.getlength(texte) / 1000 * pt / 72
    except ImportError:
        pass
    return len(texte) * 0.53 * pt / 72


def _largeur_consolas(texte, pt):
    return len(texte) * 0.55 * pt / 72


# --- intercalaires ----------------------------------------------------------
_ORDRE_RPR = ["a:ln", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill",
              "a:grpFill", "a:effectLst", "a:effectDag", "a:highlight", "a:uLnTx", "a:uLn",
              "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick",
              "a:hlinkMouseOver", "a:rtl", "a:extLst"]


def _reordonner(rpr):
    enfants = list(rpr)
    for e in enfants:
        rpr.remove(e)
    rang = {qn(t): i for i, t in enumerate(_ORDRE_RPR)}
    for e in sorted(enfants, key=lambda e: rang.get(e.tag, len(rang))):
        rpr.append(e)


def _poser_police(rpr, police):
    for tag in ("a:latin", "a:cs"):
        e = rpr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rpr, qn(tag))
        for att in list(e.attrib):
            del e.attrib[att]
        e.set("typeface", police)


def uniformiser_titre_intercalaire(boite, largeur_diapo):
    """Impose le style commun au titre d'un intercalaire. Rend la taille posee."""
    tf = boite.text_frame
    paragraphes = [p for p in tf.paragraphs if p.runs]
    disponible = (largeur_diapo - boite.left - MARGE_DROITE) / 914400
    body = tf._txBody.find(qn("a:bodyPr"))
    marges = (int(body.get("lIns", 91440)) + int(body.get("rIns", 91440))) / 914400

    def est_code(r):
        return (r.font.name or "") == "Consolas"

    def largeur(pt):
        return max(sum(_largeur_consolas(r.text, cran_dessous(pt)) if est_code(r)
                       else _largeur_calibri_gras(r.text, pt) for r in p.runs)
                   for p in paragraphes) + marges + 0.1

    taille = TAILLE_INTERCALAIRE
    while largeur(taille) > disponible and taille > TAILLE_INTERCALAIRE_MIN:
        taille = cran_dessous(taille)

    for p in tf.paragraphs:
        ppr = p._p.find(qn("a:pPr"))
        if ppr is not None and "algn" in ppr.attrib:
            del ppr.attrib["algn"]
        fin = p._p.find(qn("a:endParaRPr"))
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()
            # l'ombre et le contour vivent parfois seulement dans endParaRPr
            if fin is not None:
                for tag in ("a:ln", "a:effectLst"):
                    if rpr.find(qn(tag)) is None and fin.find(qn(tag)) is not None:
                        rpr.append(etree.fromstring(etree.tostring(fin.find(qn(tag)))))
            code = est_code(r)
            rpr.set("b", "1")
            rpr.set("sz", str(int(round((cran_dessous(taille) if code else taille) * 100))))
            if not code:
                _poser_police(rpr, POLICE_INTERCALAIRE)
            _reordonner(rpr)

    boite.width = Emu(int(min(largeur(taille), disponible) * 914400))
    return taille
