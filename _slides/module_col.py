#!/usr/bin/env python3
"""Configuration du module COL pour construire_deck.py et finaliser_deck.py.

Arborescence attendue, relative au dossier de ce fichier :
    gabarit_deck.pptx         gabarit commun : theme, dispositions, 4 slides modeles
    d05.pptx                  ancien deck 05 (05__Java_IN2_2026_-_Collections.pptx
                              renomme), qui fournit l'intercalaire
    scenario/Scenario_COL.md  le scenario

Usage :
    python3 construire_deck.py module_col COL_deck.pptx
    python3 finaliser_deck.py COL COL_deck.pptx COL_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent

CODE = "COL"
ORDRE = 9    # rang du deck dans le cours : apres INI, POO, TMP, TAB, ENU, ABS, ITF, PKG
SCENARIO = ICI / "scenario" / "Scenario_COL.md"
GABARIT = ICI / "gabarit_deck.pptx"
SOURCES = {"d05": ICI / "d05.pptx"}

# Slides prelevees dans les decks sources : (deck, numero[, nom de forme])
GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d05", 4)        # « Les collections », la pile de colis
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"COL": "fascicule COL", "G": "fil rouge GameHub"}

# Lu par finaliser_deck.py.
# SECTIONS : le texte des intercalaires du deck, numero retire, mot pour mot
#            (casse mise a part) ; un intercalaire absent de cette liste
#            n'est ni numerote ni repere.
# LIBELLES : la version courte, pour le sommaire et le repere en bas de slide.
SECTIONS = [
    "Une structure qui grandit",
    "Une liste, des positions",
    "Un ensemble, sans doublons",
    "Une table, des clés",
    "Passer un comportement",
    "Des collections emboîtées",
    "Une file, deux bouts",
]

LIBELLES = [
    "Une structure qui grandit",
    "Les listes",
    "Les ensembles",
    "Les tables associatives",
    "Passer un comportement",
    "Collections emboîtées",
    "Les files à deux bouts",
]

# Titre des intercalaires, sans numero : le finaliseur le pose.
# Les titres sont ceux de SECTIONS, dans le meme ordre : ils en sont tires,
# pour qu'un titre change a un seul endroit.
INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}

# Une seule image pour les sept sections : celle de l'ancien deck, deja
# dans la slide modele. Frederic les differenciera ensuite dans PowerPoint.
IMAGES_INTERCALAIRE = {}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "demo", "Demo_COL"),
    ("1.3", "new", None), ("1.4", "new", None), ("1.5", "new", None),

    ("2.1", "inter", None),
    ("2.2", "new", None), ("2.3", "new", None), ("2.4", "new", None),
    ("2.5", "new", None), ("2.6", "new", None), ("2.7", "renvoi", None),

    ("3.1", "inter", None),
    ("3.2", "new", None), ("3.3", "new", None), ("3.4", "table", None),
    ("3.5", "renvoi", None),

    ("4.1", "inter", None),
    ("4.2", "new", None), ("4.3", "new", None), ("4.4", "new", None),
    ("4.5", "renvoi", None),

    ("5.1", "inter", None),
    ("5.2", "new", None), ("5.3", "new", None), ("5.4", "new", None),
    ("5.5", "renvoi", None),

    ("6.1", "inter", None),
    ("6.2", "new", None), ("6.3", "new", None), ("6.4", "renvoi", None),

    ("7.1", "inter", None),
    ("7.2", "new", None), ("7.3", "new", None),

    ("8.1", "renvoi", None), ("8.2", "renvoi", None),
]


# L'intercalaire du deck 05 porte son titre en 60 pt : les titres de ce
# module, numero compris (« 6. Des collections emboitees », 28 caracteres),
# passeraient a la ligne et mordraient sur l'image. On les pose tous a la
# meme taille reduite, comme ITF (44) et PKG (54).
TAILLE_INTERCALAIRE = 48


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
