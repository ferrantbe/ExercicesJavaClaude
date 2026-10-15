#!/usr/bin/env python3
"""Configuration du module LAM pour construire_deck.py et finaliser_deck.py.

Arborescence attendue, relative au dossier de ce fichier :
    sources/gabarit_deck.pptx  gabarit commun : theme, dispositions, 4 slides modeles
    sources/d08.pptx           ancien deck 08 (generiques et lambdas reunis), qui
                               fournit l'intercalaire ; appele d08a dans SOURCES,
                               nom historique conserve
    scenario/Scenario_LAM.md  le scenario

Usage :
    python3 construire_deck.py module_lam LAM_deck.pptx
    python3 finaliser_deck.py LAM LAM_deck.pptx LAM_slides.pptx
"""
from pathlib import Path

ICI = Path(__file__).resolve().parent

CODE = "LAM"
ORDRE = 13   # rang du deck dans le cours : juste apres GEN
SCENARIO = ICI / "scenario" / "Scenario_LAM.md"
GABARIT = ICI / "sources" / "gabarit_deck.pptx"
SOURCES = {"d08a": ICI / "sources" / "d08.pptx"}

# Slides prelevees dans les decks sources : (deck, numero[, nom de forme])
GABARIT_TITRE = ("gabarit", 1)
GABARIT_INTERCALAIRE = ("d08a", 56)      # « Les expressions Lambda », pleine page
ICONE_DEMO = ("gabarit", 3, "Picture 3")
IMAGE_EXO = ("gabarit", 4, "Picture 12")

FASCICULES = {"LAM": "fascicule LAM", "G": "fil rouge GameHub"}

# Lu par finaliser_deck.py.
# SECTIONS : le texte des intercalaires du deck, numero retire, mot pour mot
#            (casse mise a part) ; un intercalaire absent de cette liste
#            n'est ni numerote ni repere.
# LIBELLES : la version courte, pour le sommaire et le repere en bas de slide.
SECTIONS = [
    "Une classe qui perd son nom",
    "Ce que le compilateur sait déjà",
    "Réutiliser ce qui existe",
]

LIBELLES = [
    "La classe anonyme",
    "La lambda",
    "JDK et références",
]

# Titre des intercalaires, sans numero : le finaliseur le pose.
# Les titres sont ceux de SECTIONS, dans le meme ordre : ils en sont tires,
# pour qu'un titre change a un seul endroit. Pas de sous-titre : celui de
# l'ancien deck (« Une classe anonyme, c'est long... ») est retire.
INTERCALAIRES = {f"{n}.1": (titre, None) for n, titre in enumerate(SECTIONS, 1)}

# Une seule image pour les trois sections : celle de l'ancien deck, deja
# dans la slide modele. Frederic les differenciera ensuite dans PowerPoint.
IMAGES_INTERCALAIRE = {}

PLAN = [
    ("0.1", "titre", None),

    ("1.1", "inter", None),
    ("1.2", "new", None), ("1.3", "new", None),
    ("1.4", "demo", "Demo_LAM"),

    ("2.1", "inter", None),
    ("2.2", "new", None), ("2.2a", "new", None), ("2.3", "new", None), ("2.4", "new", None),
    ("2.5", "renvoi", None),

    ("3.1", "inter", None),
    ("3.2", "table", None), ("3.3", "new", None), ("3.4", "new", None),
    ("3.5", "new", None), ("3.5a", "table", None), ("3.6", "renvoi", None),

    ("4.1", "new", None), ("4.2", "renvoi", None),
]


# Meme valeur qu'EXC et ITF : le plus long titre de ce module, numero compris
# (« 2. Ce que le compilateur sait deja », 34 caracteres), passerait a la
# ligne a la taille de l'ancien deck. A confirmer dans PowerPoint : Corbel
# n'est pas disponible pour le rendu de controle.
TAILLE_INTERCALAIRE = 44


def retoucher(slide, ident, bloc):
    if ident not in INTERCALAIRES or TAILLE_INTERCALAIRE is None:
        return
    from pptx.util import Pt
    boites = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.strip() and not s.is_placeholder]
    if not boites:
        return
    titre = min(boites, key=lambda s: (s.top, s.left))
    for par in titre.text_frame.paragraphs:
        for run in par.runs:
            run.font.size = Pt(TAILLE_INTERCALAIRE)
